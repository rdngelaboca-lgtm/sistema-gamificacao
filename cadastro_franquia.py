# -*- coding: utf-8 -*-
# ==============================================================================
# == cadastro_franquia.py - Pedido de cadastro de produtos na Franquia ===========
# ==============================================================================
# Produto novo na loja precisa ser cadastrado na Franquia antes de vender. Para cada
# produto a franquia pede: nome (limpo), código de barras, NCM, preço de custo e preço
# de venda. Aqui ficam as regras (o programa do Estoque, aba "🏷️ Cadastro Franquia", só
# mostra e chama estas funções):
#   - CUSTO = preço da NOTA por unidade do estoque (com impostos e frete, SEM royalties):
#     a nota nova que ainda está chegando (XML da SEFAZ) ou a última compra paga;
#   - PREÇO DE VENDA = custo × MARKUP (multiplicador: 2,5 = custo × 2,5), arredondado
#     para cima terminando em ,90 (R$ 24,37 -> R$ 24,90). Markup padrão por categoria;
#   - CÓDIGO DE BARRAS = o da UNIDADE (numa caixa de 12, o XML traz o da caixa em cEAN e
#     o da unidade em cEANTrib);
#   - situação de cada produto: para enviar / enviado (aguardando a franquia) /
#     cadastrado / não vende (insumos, embalagens...);
#   - planilha do Excel com as 5 colunas para mandar à franquia.
# ==============================================================================
import logging
import os
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_FLOOR, ROUND_HALF_UP

import database
import nota_xml
from texto_produto import sugerir_nome_limpo

logger = logging.getLogger(__name__)

PENDENTE, ENVIADO, CADASTRADO, NAO_VENDE = 'pendente', 'enviado', 'cadastrado', 'nao_vende'
STATUS = (PENDENTE, ENVIADO, CADASTRADO, NAO_VENDE)
TEXTO_STATUS = {PENDENTE: '📝 Para enviar', ENVIADO: '📤 Enviado (aguardando)', CADASTRADO: '✅ Cadastrado',
                NAO_VENDE: '🚫 Não vende'}
COLUNAS_PLANILHA = ('Nome do produto', 'Código de barras', 'NCM', 'Preço de custo', 'Preço de venda')
MARKUP_MAXIMO = Decimal('100')
_tabelas_ok = False


class ErroCadastro(Exception):
    """Erro com mensagem para mostrar ao usuário."""


# ------------------------------------------------------------------------------
# Números
# ------------------------------------------------------------------------------
def numero(valor, nome='Valor', vazio=None):
    """'2,5' / '2.5' / 'R$ 1.234,56' / '2,5x' -> Decimal. Vazio -> 'vazio'. Inválido -> ErroCadastro."""
    if valor is None:
        return vazio
    if isinstance(valor, (int, Decimal)):
        return Decimal(valor)
    if isinstance(valor, float):
        return Decimal(str(valor))
    t = str(valor).replace('R$', '').replace(' ', '').strip().rstrip('xX×')
    if not t:
        return vazio
    if ',' in t and '.' in t:
        if t.rfind(',') > t.rfind('.'):
            t = t.replace('.', '').replace(',', '.')
        else:
            t = t.replace(',', '')
    else:
        t = t.replace(',', '.')
    try:
        d = Decimal(t)
    except InvalidOperation:
        raise ErroCadastro(f"{nome} inválido: '{valor}'.")
    if not d.is_finite():
        raise ErroCadastro(f"{nome} inválido: '{valor}'.")
    return d


def arredondar_90(valor):
    """Preço de venda terminando em ,90, SEMPRE para cima: 24,37 -> 24,90; 24,95 -> 25,90; 24,90 -> 24,90."""
    if valor is None:
        return None
    v = Decimal(valor).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    if v <= 0:
        return None
    base = v.to_integral_value(rounding=ROUND_FLOOR)
    alvo = base + Decimal('0.90')
    return alvo if v <= alvo else alvo + 1


def preco_venda_sugerido(custo, markup):
    """custo × markup (multiplicador), arredondado em ,90. None se faltar custo ou markup."""
    if custo is None or markup is None or Decimal(custo) <= 0 or Decimal(markup) <= 0:
        return None
    return arredondar_90(Decimal(custo) * Decimal(markup))


def markup_de(custo, venda):
    """Markup que um preço de venda representa (venda ÷ custo), com 2 casas. None se não dá para calcular."""
    if custo is None or venda is None or Decimal(custo) <= 0:
        return None
    return (Decimal(venda) / Decimal(custo)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def ean_valido(ean):
    return database.ean_valido(ean)


def ncm_valido(ncm):
    return database.ncm_valido(ncm)


# ------------------------------------------------------------------------------
# Banco
# ------------------------------------------------------------------------------
def _conexao():
    conn = database.get_db_connection()
    if not conn:
        raise ErroCadastro("Sem conexão com o banco de dados.")
    return conn


def garantir_tabelas():
    """Cria (uma vez) a tabela CadastroFranquia e o markup padrão das categorias."""
    global _tabelas_ok
    if _tabelas_ok:
        return
    database._garantir_colunas_estoque()          # EANUnidade nos vínculos
    conn = _conexao()
    try:
        cur = conn.cursor()
        cur.execute("""
            IF NOT EXISTS (SELECT * FROM sys.tables WHERE name = 'CadastroFranquia')
            CREATE TABLE CadastroFranquia (
                ProdutoID INT NOT NULL PRIMARY KEY,
                Status VARCHAR(20) NOT NULL,
                NomeCadastro NVARCHAR(150) NULL,
                EAN VARCHAR(14) NULL,
                NCM VARCHAR(10) NULL,
                PrecoCusto DECIMAL(18, 4) NULL,
                Markup DECIMAL(9, 4) NULL,
                PrecoVenda DECIMAL(18, 2) NULL,
                EnviadoEm DATETIME NULL,
                CadastradoEm DATETIME NULL,
                AtualizadoEm DATETIME NULL
            )
        """)
        cur.execute("IF COL_LENGTH('CategoriasProduto', 'MarkupPadrao') IS NULL "
                    "ALTER TABLE CategoriasProduto ADD MarkupPadrao DECIMAL(9, 4) NULL")
        conn.commit()
        _tabelas_ok = True
    except Exception as e:
        conn.rollback()
        logger.error(f"Cadastro franquia: não deu para preparar o banco: {e}", exc_info=True)
        raise ErroCadastro(f"Não deu para preparar o banco: {e}")
    finally:
        conn.close()


def markups_por_categoria():
    """{categoria: markup padrão} (só as que têm)."""
    garantir_tabelas()
    conn = _conexao()
    try:
        cur = conn.cursor()
        cur.execute("SELECT NomeCategoria, MarkupPadrao FROM CategoriasProduto WHERE MarkupPadrao IS NOT NULL")
        return {r[0]: Decimal(str(r[1])) for r in cur.fetchall() if r[1] is not None and Decimal(str(r[1])) > 0}
    finally:
        conn.close()


def definir_markup_categoria(categoria, texto):
    """Markup padrão da categoria ('2,5'). Vazio ou 0 = sem padrão. Devolve a mensagem."""
    m = numero(texto, 'Markup')
    if m is not None and (m < 0 or m > MARKUP_MAXIMO):
        raise ErroCadastro(f"Markup precisa estar entre 0 e {MARKUP_MAXIMO} (ex.: 2,5 = custo × 2,5).")
    garantir_tabelas()
    conn = _conexao()
    try:
        cur = conn.cursor()
        cur.execute("UPDATE CategoriasProduto SET MarkupPadrao = ? WHERE NomeCategoria = ?",
                    (m if m else None), categoria)
        if getattr(cur, 'rowcount', 1) == 0:
            raise ErroCadastro(f"Categoria '{categoria}' não encontrada.")
        conn.commit()
    finally:
        conn.close()
    return (f"Markup padrão de {categoria}: {str(m).replace('.', ',')}" if m else f"{categoria} ficou sem markup padrão.")


# ------------------------------------------------------------------------------
# Notas que estão chegando (XML da SEFAZ ainda não lançado)
# ------------------------------------------------------------------------------
def _so_digitos(texto):
    return ''.join(ch for ch in str(texto or '') if ch.isdigit())


class _Reconhecedor:
    """
    O mesmo reconhecimento de database.buscar_vinculo_inteligente (descrição exata; depois cProd;
    depois EAN, só se não for ambíguo), mas com os vínculos carregados de UMA vez: ler 20 notas de
    40 itens não pode abrir 800 conexões no banco (a aba ficaria lenta no SQL Server de verdade).
    """
    def __init__(self, cur):
        cur.execute("SELECT FornecedorID, CNPJ FROM Fornecedores ORDER BY FornecedorID")
        self.fornecedores = {}
        for fid, cnpj in cur.fetchall():
            if _so_digitos(cnpj):
                self.fornecedores.setdefault(_so_digitos(cnpj), fid)
        cur.execute("SELECT ProdutoFornecedorID, FornecedorID, ProdutoID, FatorConversao, DescricaoXML, CodigoFornecedor, EAN "
                    "FROM ProdutosFornecedor ORDER BY ProdutoFornecedorID")
        self.por_desc, self.por_cod, self.por_ean = {}, {}, {}
        for pf, fid, pid, fator, desc, cod, ean in cur.fetchall():
            v = {'ProdutoFornecedorID': pf, 'ProdutoID': pid, 'Fator': fator}
            self.por_desc.setdefault((fid, self._desc(desc)), v)
            if pid is None:
                continue
            if database._codigo_valido(cod):
                self.por_cod.setdefault((fid, database._codigo_valido(cod)), []).append(v)
            if ean_valido(ean):
                self.por_ean.setdefault((fid, ean_valido(ean)), []).append(v)

    @staticmethod
    def _desc(texto):
        return str(texto or '').rstrip().casefold()     # o SQL Server compara sem maiúsculas e sem espaço no fim

    def fornecedor(self, cnpj):
        return self.fornecedores.get(_so_digitos(cnpj))

    def vinculo(self, fornecedor_id, descricao, cprod=None, ean=None):
        if not fornecedor_id:
            return None
        v = self.por_desc.get((fornecedor_id, self._desc(descricao)))
        if v:
            return v
        for mapa, chave in ((self.por_cod, database._codigo_valido(cprod)), (self.por_ean, ean_valido(ean))):
            achados = mapa.get((fornecedor_id, chave)) if chave else None
            if not achados:
                continue
            fatores = {Decimal(str(a['Fator'])) if a['Fator'] is not None else Decimal('1') for a in achados}
            if len({a['ProdutoID'] for a in achados}) == 1 and len(fatores) == 1:
                return achados[-1]                       # o vínculo mais novo
        return None


def _reconhecedor():
    conn = _conexao()
    try:
        return _Reconhecedor(conn.cursor())
    finally:
        conn.close()


def _pasta_sefaz():
    import nfe_distribuicao as nd
    return nd.pasta_xml()


def itens_das_notas_novas(pasta=None):
    """
    Lê os XMLs da pasta da SEFAZ (as notas que ainda não foram lançadas no estoque):
    {'por_produto': {ProdutoID: {'custo', 'ean_unidade', 'ncm', 'nf', 'fornecedor', 'data', 'descricao'}},
     'sem_produto': [{'nf', 'fornecedor', 'data', 'descricao', 'ean', 'ncm', 'custo', 'arquivo'}]}
    'custo' = preço da nota por unidade do estoque (preço do item ÷ Qtd/Cx do vínculo).
    Item sem vínculo = produto NOVO: precisa ser vinculado/criado na aba 3 antes de ir para o cadastro.
    """
    pasta = pasta or _pasta_sefaz()
    por_produto, sem_produto = {}, []
    if not pasta or not os.path.isdir(pasta):
        return {'por_produto': por_produto, 'sem_produto': sem_produto}
    arquivos = sorted(f for f in os.listdir(pasta) if f.lower().endswith('.xml'))
    rec = _reconhecedor() if arquivos else None
    for arquivo in arquivos:
        try:
            cab, itens = nota_xml.ler_xml_nota_fiscal(os.path.join(pasta, arquivo))
        except Exception as e:
            logger.warning(f"Cadastro franquia: XML {arquivo} não abriu: {e}")
            continue
        if cab.get('Finalidade') == '4':
            continue
        forn_id = rec.fornecedor(cab.get('FornecedorCNPJ'))
        data_nota = database._como_data(cab.get('DataEmissao'))
        for it in itens:
            tipo = nota_xml.tipo_item_por_cfop(it.get('CFOP'))
            if tipo == 'ignorar':
                continue
            v = rec.vinculo(forn_id, it['DescricaoXML'], it.get('cProd'), it.get('cEAN'))
            fator = Decimal(str(v['Fator'])) if v and v.get('Fator') and Decimal(str(v['Fator'])) > 0 else Decimal('1')
            custo = (it['PrecoCustoUnitario'] / fator) if tipo == 'compra' and it['PrecoCustoUnitario'] > 0 else None
            ean_caixa, ean_trib = ean_valido(it.get('cEAN')), ean_valido(it.get('cEANTrib'))
            ean_unidade = ean_trib if (ean_trib and (fator > 1 or ean_trib != ean_caixa)) else (ean_caixa if fator == 1 else None)
            dados = {'custo': custo, 'ean_unidade': ean_unidade, 'ncm': ncm_valido(it.get('NCM')),
                     'nf': str(cab.get('NumeroNF') or ''), 'fornecedor': cab.get('FornecedorNome') or '', 'data': data_nota,
                     'descricao': it['DescricaoXML'], 'arquivo': arquivo}
            if not v or not v.get('ProdutoID'):
                sem_produto.append(dict(dados, ean=ean_trib or ean_caixa, custo=it['PrecoCustoUnitario']
                                        if tipo == 'compra' else None))
                continue
            atual = por_produto.get(v['ProdutoID'])
            if atual is None or (data_nota or date.min) >= (atual['data'] or date.min):
                if atual and custo is None:          # bonificação não apaga o custo de outra nota
                    dados['custo'] = atual['custo']
                por_produto[v['ProdutoID']] = dados
    return {'por_produto': por_produto, 'sem_produto': sem_produto}


def produtos_novos_nas_notas(chaves, pasta=None):
    """[AVISO TELEGRAM] Quantos itens destas notas ainda não têm produto no estoque (produtos novos)."""
    pasta = pasta or _pasta_sefaz()
    novos = 0
    rec = _reconhecedor() if chaves else None
    for chave in chaves or []:
        caminho = os.path.join(pasta, f"{chave}.xml")
        if not os.path.isfile(caminho):
            continue
        try:
            cab, itens = nota_xml.ler_xml_nota_fiscal(caminho)
        except Exception:
            continue
        forn_id = rec.fornecedor(cab.get('FornecedorCNPJ'))
        for it in itens:
            if nota_xml.tipo_item_por_cfop(it.get('CFOP')) == 'ignorar':
                continue
            v = rec.vinculo(forn_id, it['DescricaoXML'], it.get('cProd'), it.get('cEAN'))
            if not v or not v.get('ProdutoID'):
                novos += 1
    return novos


# ------------------------------------------------------------------------------
# A lista
# ------------------------------------------------------------------------------
def _codigos_por_produto(cur):
    """Candidatos a código de barras da UNIDADE de cada produto: {pid: [(prioridade, ordem, ean, origem)]}."""
    cand = {}
    if database._tabela_existe(cur, 'CompraCodigos'):
        cur.execute("SELECT Codigo, ProdutoID, Fator FROM CompraCodigos")
        for cod, pid, fator in cur.fetchall():
            e = ean_valido(cod)
            if e and Decimal(str(fator or 1)) == 1:
                cand.setdefault(pid, []).append((1, 0, e, 'bipado no app'))
    cur.execute("SELECT ProdutoFornecedorID, ProdutoID, EAN, EANUnidade, FatorConversao FROM ProdutosFornecedor "
                "WHERE ProdutoID IS NOT NULL")
    for pf, pid, ean, ean_unid, fator in cur.fetchall():
        f = Decimal(str(fator)) if fator is not None and Decimal(str(fator)) > 0 else Decimal('1')
        e, eu = ean_valido(ean), ean_valido(ean_unid)
        if e and f == 1:
            cand.setdefault(pid, []).append((2, pf, e, 'nota do fornecedor'))
        if eu:
            cand.setdefault(pid, []).append((3, pf, eu, 'nota (unidade da caixa)'))
        if e and f > 1:
            cand.setdefault(pid, []).append((5, pf, e, 'CAIXA'))
    return cand


def _melhor_codigo(candidatos, da_nota_nova=None):
    if da_nota_nova:
        candidatos = list(candidatos or []) + [(4, 0, da_nota_nova, 'nota que está chegando')]
    if not candidatos:
        return None, None
    prioridade = min(c[0] for c in candidatos)
    escolhido = max((c for c in candidatos if c[0] == prioridade), key=lambda c: c[1])
    return escolhido[2], escolhido[3]


def _custos_manuais(cur):
    cur.execute("""
        SELECT PF.ProdutoID, INI.PrecoCustoUnitario, NF.DataEmissao, NF.NotaID
        FROM ItensNotaFiscalEntrada INI
        JOIN NotasFiscaisEntrada NF ON INI.NotaID = NF.NotaID
        JOIN ProdutosFornecedor PF ON INI.ProdutoFornecedorID = PF.ProdutoFornecedorID
        JOIN Fornecedores F ON NF.FornecedorID = F.FornecedorID
        WHERE PF.ProdutoID IS NOT NULL AND F.CNPJ = ?
    """, database.CNPJ_FORNECEDOR_INTERNO)
    manuais = {}
    for pid, custo, dt, nota_id in cur.fetchall():
        chave = (database._como_data(dt) or date.min, nota_id or 0)
        if Decimal(str(custo or 0)) > 0 and (pid not in manuais or chave > manuais[pid][0]):
            manuais[pid] = (chave, Decimal(str(custo)))
    return {pid: v[1] for pid, v in manuais.items()}


def _dec_ou_none(v):
    return Decimal(str(v)) if v is not None else None


def listar(incluir_notas=True, pasta=None):
    """
    Uma linha por produto do Catálogo:
    {'ProdutoID', 'Produto' (nome no Catálogo), 'Unidade', 'Categoria', 'Status', 'TextoStatus',
     'Nome', 'EAN', 'OrigemEAN', 'NCM', 'Custo', 'OrigemCusto', 'Markup', 'MarkupPadrao', 'PrecoVenda',
     'MarkupReal', 'Chegando' (dados da nota nova ou None), 'EnviadoEm', 'CadastradoEm', 'Avisos': [...]}
    Produto ENVIADO ou CADASTRADO mostra os valores que foram mandados (não muda com notas novas).
    Para enviar: o que foi digitado e salvo; no resto, a sugestão do sistema.
    """
    garantir_tabelas()
    notas = itens_das_notas_novas(pasta) if incluir_notas else {'por_produto': {}, 'sem_produto': []}
    precos = database.ultimos_precos_pagos()
    conn = _conexao()
    try:
        cur = conn.cursor()
        cur.execute("SELECT ProdutoID, NomeProduto, UnidadeMedida, Categoria, NCM FROM ProdutosEstoque")
        produtos = cur.fetchall()
        inativos = database.ids_produtos_inativos(cur)     # [PRODUTO INATIVO]
        cur.execute("SELECT ProdutoID, Status, NomeCadastro, EAN, NCM, PrecoCusto, Markup, PrecoVenda, EnviadoEm, "
                    "CadastradoEm FROM CadastroFranquia")
        salvos = {r[0]: r for r in cur.fetchall()}
        cur.execute("SELECT NomeCategoria, MarkupPadrao FROM CategoriasProduto")
        markups = {r[0]: Decimal(str(r[1])) for r in cur.fetchall() if r[1] is not None and Decimal(str(r[1])) > 0}
        codigos = _codigos_por_produto(cur)
        cur.execute("SELECT ProdutoID, NCM FROM ProdutosFornecedor WHERE ProdutoID IS NOT NULL ORDER BY ProdutoFornecedorID")
        ncm_vinculo = {}
        for pid, ncm in cur.fetchall():
            if ncm_valido(ncm):
                ncm_vinculo[pid] = ncm_valido(ncm)
        manuais = _custos_manuais(cur)
    finally:
        conn.close()

    linhas = []
    for pid, nome_cat, un, categoria, ncm_prod in produtos:
        s = salvos.get(pid)
        status = (s[1] if s else PENDENTE) or PENDENTE
        chegando = notas['por_produto'].get(pid)
        if pid in inativos and not chegando:      # [PRODUTO INATIVO] só aparece se está chegando numa nota nova
            continue
        # --- sugestões do sistema ---
        ean_sug, origem_ean = _melhor_codigo(codigos.get(pid), chegando['ean_unidade'] if chegando else None)
        ncm_sug = ncm_valido(ncm_prod) or ncm_vinculo.get(pid) or (chegando['ncm'] if chegando else None)
        ultimo = precos.get(pid)
        if chegando and chegando.get('custo') and (not ultimo or (chegando['data'] or date.min) >= (ultimo['data'] or date.min)):
            custo_sug = chegando['custo']
            origem_custo = (f"nota {chegando['nf']} de {chegando['data']:%d/%m/%Y} (chegando)" if chegando.get('data')
                            else f"nota {chegando['nf']} (chegando)")
        elif ultimo:
            custo_sug = ultimo['preco']
            origem_custo = (f"última compra {ultimo['data']:%d/%m/%Y} ({ultimo['fornecedor']})" if ultimo.get('data')
                            else f"última compra ({ultimo['fornecedor']})")
        elif pid in manuais:
            custo_sug, origem_custo = manuais[pid], "custo manual do Catálogo"
        else:
            custo_sug, origem_custo = None, "sem compra"
        markup_cat = markups.get(categoria or 'Geral')
        # --- o que vale ---
        nome = (s[2] if s and s[2] else None) or sugerir_nome_limpo(nome_cat)
        ean = (s[3] if s and s[3] else None) or ean_sug
        if s and s[3]:
            origem_ean = 'digitado'
        ncm = (s[4] if s and s[4] else None) or ncm_sug
        custo = _dec_ou_none(s[5]) if s and s[5] is not None else custo_sug
        if s and s[5] is not None:
            origem_custo = 'enviado à franquia' if status in (ENVIADO, CADASTRADO) else 'digitado'
        markup = _dec_ou_none(s[6]) if s and s[6] is not None else markup_cat
        venda = _dec_ou_none(s[7]) if s and s[7] is not None else preco_venda_sugerido(custo, markup)
        avisos = []
        if not ean:
            avisos.append('sem código de barras')
        elif origem_ean == 'CAIXA':
            avisos.append('código da CAIXA (confira o da unidade)')
        if not ncm_valido(ncm):
            avisos.append('sem NCM')
        if not custo:
            avisos.append('sem custo')
        if not venda:
            avisos.append('sem preço de venda')
        elif custo and venda <= custo:
            avisos.append('venda menor que o custo')
        linhas.append({
            'ProdutoID': pid, 'Produto': nome_cat or f'Produto {pid}', 'Unidade': (un or 'UN').strip() or 'UN',
            'Categoria': categoria or 'Geral', 'Status': status, 'TextoStatus': TEXTO_STATUS.get(status, status),
            'Nome': nome, 'EAN': ean, 'OrigemEAN': origem_ean, 'NCM': ncm, 'Custo': custo, 'OrigemCusto': origem_custo,
            'Markup': markup, 'MarkupPadrao': markup_cat, 'PrecoVenda': venda, 'MarkupReal': markup_de(custo, venda),
            'Chegando': chegando, 'EnviadoEm': s[8] if s else None, 'CadastradoEm': s[9] if s else None,
            'Avisos': avisos})
    linhas.sort(key=lambda l: (database.sem_acento_simples(l['Nome'])))
    return {'linhas': linhas, 'sem_produto': notas['sem_produto']}


# ------------------------------------------------------------------------------
# Gravar
# ------------------------------------------------------------------------------
def _upsert(cur, pid, campos):
    cur.execute("SELECT Status FROM CadastroFranquia WHERE ProdutoID = ?", int(pid))
    existe = cur.fetchone()
    campos = dict(campos, AtualizadoEm=datetime.now())
    if existe:
        cur.execute(f"UPDATE CadastroFranquia SET {', '.join(f'{c} = ?' for c in campos)} WHERE ProdutoID = ?",
                    *campos.values(), int(pid))
    else:
        campos.setdefault('Status', PENDENTE)
        cur.execute(f"INSERT INTO CadastroFranquia (ProdutoID, {', '.join(campos)}) VALUES (?, {', '.join('?' for _ in campos)})",
                    int(pid), *campos.values())


def salvar(produto_id, nome=None, ean=None, ncm=None, custo=None, markup=None, venda=None):
    """
    Grava o que o gestor digitou para um produto (campo None = não mexe; '' = volta para a sugestão).
    Valida: código de barras 8 a 14 números, NCM 8 números, custo/venda > 0, markup entre 0 e 100.
    """
    campos = {}
    if nome is not None:
        n = ' '.join(str(nome).split())
        if len(n) > 150:
            raise ErroCadastro("Nome muito comprido (máximo 150 letras).")
        campos['NomeCadastro'] = n or None
    if ean is not None:
        t = str(ean).strip()
        if t and not ean_valido(t):
            raise ErroCadastro(f"Código de barras inválido: '{ean}' (8 a 14 números).")
        campos['EAN'] = ean_valido(t) if t else None
    if ncm is not None:
        t = str(ncm).strip()
        if t and not ncm_valido(t):
            raise ErroCadastro(f"NCM inválido: '{ncm}' (8 números).")
        campos['NCM'] = ncm_valido(t) if t else None
    if custo is not None:
        c = numero(custo, 'Custo')
        if c is not None and c <= 0:
            raise ErroCadastro("O custo precisa ser maior que zero.")
        campos['PrecoCusto'] = c.quantize(Decimal('0.0001')) if c is not None else None
    if markup is not None:
        m = numero(markup, 'Markup')
        if m is not None and (m <= 0 or m > MARKUP_MAXIMO):
            raise ErroCadastro(f"Markup precisa ser maior que 0 e até {MARKUP_MAXIMO} (ex.: 2,5 = custo × 2,5).")
        campos['Markup'] = m.quantize(Decimal('0.0001')) if m is not None else None
    if venda is not None:
        v = numero(venda, 'Preço de venda')
        if v is not None and v <= 0:
            raise ErroCadastro("O preço de venda precisa ser maior que zero.")
        campos['PrecoVenda'] = v.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP) if v is not None else None
    if not campos:
        return "Nada para gravar."
    garantir_tabelas()
    conn = _conexao()
    try:
        cur = conn.cursor()
        _upsert(cur, produto_id, campos)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return "Gravado."


def aplicar_markup(produto_ids, markup):
    """Mesmo markup para vários produtos (o preço de venda é recalculado). Enviados/cadastrados não mudam."""
    m = numero(markup, 'Markup')
    if m is None or m <= 0 or m > MARKUP_MAXIMO:
        raise ErroCadastro(f"Markup precisa ser maior que 0 e até {MARKUP_MAXIMO} (ex.: 2,5).")
    garantir_tabelas()
    conn = _conexao()
    feitos, pulados = 0, 0
    try:
        cur = conn.cursor()
        for pid in produto_ids:
            cur.execute("SELECT Status FROM CadastroFranquia WHERE ProdutoID = ?", int(pid))
            r = cur.fetchone()
            if r and r[0] in (ENVIADO, CADASTRADO):
                pulados += 1
                continue
            _upsert(cur, pid, {'Markup': m.quantize(Decimal('0.0001')), 'PrecoVenda': None})
            feitos += 1
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return (f"Markup {str(m).replace('.', ',')} em {feitos} produto(s)."
            + (f" {pulados} já enviado(s)/cadastrado(s) não mudaram." if pulados else ""))


def voltar_para_sugestao(produto_ids):
    """Apaga custo, markup e preço digitados: volta a valer o que o sistema sugere (custo da nota mais nova)."""
    garantir_tabelas()
    conn = _conexao()
    try:
        cur = conn.cursor()
        for pid in produto_ids:
            cur.execute("UPDATE CadastroFranquia SET PrecoCusto = NULL, Markup = NULL, PrecoVenda = NULL, AtualizadoEm = ? "
                        "WHERE ProdutoID = ? AND Status = ?", datetime.now(), int(pid), PENDENTE)
        conn.commit()
    finally:
        conn.close()


def marcar(linhas_ou_ids, status):
    """
    Muda a situação. 'linhas_ou_ids' = linhas de listar() (para ENVIADO/CADASTRADO os valores mostrados
    ficam GRAVADOS: é o que foi mandado à franquia) ou só os ProdutoID.
    """
    if status not in STATUS:
        raise ErroCadastro("Situação inválida.")
    garantir_tabelas()
    agora = datetime.now()
    conn = _conexao()
    try:
        cur = conn.cursor()
        for item in linhas_ou_ids:
            linha = item if isinstance(item, dict) else None
            pid = linha['ProdutoID'] if linha else int(item)
            campos = {'Status': status}
            if status in (ENVIADO, CADASTRADO) and linha:
                campos.update({'NomeCadastro': linha.get('Nome') or None, 'EAN': ean_valido(linha.get('EAN')),
                               'NCM': ncm_valido(linha.get('NCM')),
                               'PrecoCusto': linha['Custo'].quantize(Decimal('0.0001')) if linha.get('Custo') else None,
                               'Markup': linha['Markup'].quantize(Decimal('0.0001')) if linha.get('Markup') else None,
                               'PrecoVenda': linha['PrecoVenda'] if linha.get('PrecoVenda') else None})
            if status == ENVIADO:
                campos['EnviadoEm'] = agora
            elif status == CADASTRADO:
                campos['CadastradoEm'] = agora
            elif status == PENDENTE:
                campos.update({'EnviadoEm': None, 'CadastradoEm': None})
            _upsert(cur, pid, campos)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


# ------------------------------------------------------------------------------
# Planilha para a franquia
# ------------------------------------------------------------------------------
def problemas_para_envio(linhas):
    """[(nome, problema)] do que está faltando ou estranho antes de mandar."""
    return [(l['Nome'], ', '.join(l['Avisos'])) for l in linhas if l.get('Avisos')]


def gerar_planilha(linhas, caminho):
    """Excel com as 5 colunas que a franquia pede. Código de barras e NCM vão como TEXTO (sem virar 7,89E+12)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    if not linhas:
        raise ErroCadastro("Nenhum produto para a planilha.")
    wb = Workbook()
    ws = wb.active
    ws.title = 'Cadastro'
    ws.append(list(COLUNAS_PLANILHA))
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='2E7D32')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for l in linhas:
        ws.append([l.get('Nome') or '', l.get('EAN') or '', l.get('NCM') or '',
                   float(l['Custo']) if l.get('Custo') else None, float(l['PrecoVenda']) if l.get('PrecoVenda') else None])
        linha = ws.max_row
        ws.cell(linha, 2).number_format = '@'
        ws.cell(linha, 3).number_format = '@'
        ws.cell(linha, 4).number_format = '#,##0.00'
        ws.cell(linha, 5).number_format = '#,##0.00'
    for letra, largura in zip('ABCDE', (48, 18, 12, 15, 15)):
        ws.column_dimensions[letra].width = largura
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:E{ws.max_row}"
    wb.save(caminho)
    return caminho


# ------------------------------------------------------------------------------
# Códigos de barras das notas antigas
# ------------------------------------------------------------------------------
def aprender_codigos_das_notas(pastas=None, progresso=None):
    """
    Lê os XMLs guardados (pasta da SEFAZ e 'importadas') e grava, em cada vínculo comprado em CAIXA,
    o código de barras da UNIDADE que vem no XML (cEANTrib). Para os produtos antigos.
    Devolve (arquivos lidos, códigos novos).
    """
    database._garantir_colunas_estoque()
    if pastas is None:
        base = _pasta_sefaz()
        pastas = [base, os.path.join(base, 'importadas')]
    conn = _conexao()
    try:
        cur = conn.cursor()
        rec = _Reconhecedor(cur)
        cur.execute("SELECT ProdutoFornecedorID, EANUnidade FROM ProdutosFornecedor")
        ja_tem = {pf for pf, ean_unid in cur.fetchall() if ean_valido(ean_unid)}
        arquivos = [os.path.join(p, f) for p in pastas if p and os.path.isdir(p)
                    for f in os.listdir(p) if f.lower().endswith('.xml')]
        novos = {}
        for i, caminho in enumerate(arquivos):
            if progresso and i % 25 == 0:
                progresso(i, len(arquivos))
            try:
                cab, itens = nota_xml.ler_xml_nota_fiscal(caminho)
            except Exception:
                continue
            forn_id = rec.fornecedor(cab.get('FornecedorCNPJ'))
            for it in itens:
                trib, caixa = ean_valido(it.get('cEANTrib')), ean_valido(it.get('cEAN'))
                if not trib or trib == caixa:
                    continue
                v = rec.vinculo(forn_id, it['DescricaoXML'], it.get('cProd'), it.get('cEAN'))
                if v and v.get('ProdutoID') and v['ProdutoFornecedorID'] not in ja_tem:
                    novos[v['ProdutoFornecedorID']] = trib
        for pf, trib in novos.items():
            cur.execute("UPDATE ProdutosFornecedor SET EANUnidade = ? WHERE ProdutoFornecedorID = ?", trib, pf)
        conn.commit()
        return len(arquivos), len(novos)
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
