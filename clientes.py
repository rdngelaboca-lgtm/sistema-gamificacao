# ==============================================================================
# == clientes.py - Clientes no WhatsApp (cadastro e mensagens como lista de transmissão)
# ==============================================================================
# Gestão › 📢 Clientes:
#   - cadastro (nome, WhatsApp, aniversário, etiquetas) e quem AUTORIZOU receber mensagens;
#   - mensagem com {nome} e foto (tirada no celular) para uma ou mais etiquetas, agora ou agendada;
#   - o robô (agendador.py) manda em FILA: uma mensagem a cada 6 a 15 segundos, só das 8h às 21h,
#     no máximo 300 por dia, e pausa sozinho se falhar 5 vezes seguidas (número desconectado);
#   - aniversário automático (10h), "SAIR" respondido no WhatsApp (webhook da Z-API), QR do balcão
#     (o cliente se cadastra sozinho, já autorizando) e importar planilha / agenda do celular.
# As mensagens saem por um NÚMERO SÓ DOS CLIENTES (2ª instância da Z-API, gravada no banco pela tela
# ⚙️ Ajustes; NUNCA pelo número da loja): se ele for bloqueado, a escala e os avisos continuam.
# Os tokens ficam no banco (não no config.py, que está no GitHub) e nunca vão para o log nem para a tela.
# ==============================================================================
import base64
import csv
import hmac
import io
import logging
import os
import random
import re
import secrets
import threading
import time
import unicodedata
import uuid
from datetime import date, datetime, timedelta

import config
import database

logger = logging.getLogger(__name__)

PASTA_FOTOS = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fotos_clientes')
INTERVALO_SEG = (6, 15)          # entre uma mensagem e outra (aleatório: parece gente, não robô)
LIMITE_POR_DIA = 300             # no máximo por dia, somando todos os envios
HORARIO_ENVIO = (8, 21)          # só das 8h às 21h (não incomoda à noite)
FALHAS_SEGUIDAS = 5              # 5 falhas seguidas: pausa (o número pode ter desconectado)
MAX_TEXTO = 1500
MAX_LEGENDA = 1000               # texto maior que isso com foto: vai a foto e depois o texto
MAX_FOTO_BYTES = 3 * 1024 * 1024
MAX_PLANILHA_BYTES = 5 * 1024 * 1024
MAX_LINHAS_IMPORTAR = 2000
MAX_ETIQUETAS = 10
RODAPE = "_Para não receber mais mensagens, responda SAIR._"
PALAVRAS_SAIR = {'sair', 'parar', 'pare', 'stop', 'cancelar', 'descadastrar', 'remover', 'nao quero'}
PALAVRAS_VOLTAR = {'voltar', 'quero voltar', 'entrar'}
ETIQUETA_BALCAO = 'Balcão'
CADASTROS_POR_IP = (5, 3600)     # cadastro pelo QR: no máximo 5 por hora do mesmo aparelho/rede
CADASTROS_POR_DIA = 150          # e no máximo 150 por dia no total (ninguém usa o QR para mandar boas-vindas a estranhos)
ST_AGENDADA, ST_ENVIANDO, ST_PAUSADA, ST_CONCLUIDA, ST_CANCELADA = 'agendada', 'enviando', 'pausada', 'concluida', 'cancelada'
NOMES_STATUS = {ST_AGENDADA: '⏰ Agendada', ST_ENVIANDO: '📤 Enviando', ST_PAUSADA: '⏸️ Pausada',
                ST_CONCLUIDA: '✅ Concluída', ST_CANCELADA: '⏹️ Cancelada'}
TEXTO_ANIVERSARIO_PADRAO = ("🎂 Feliz aniversário, {nome}! 🍦\n"
                            "O Gela Boca deseja um dia muito doce. Passe aqui hoje: tem um mimo esperando por você!")
TEXTO_BOAS_VINDAS = ("Oi, {nome}! 🍦 Agora você recebe as novidades e promoções do Gela Boca por aqui.\n"
                     "Salve este número na sua agenda para ver as fotos!")
TEXTO_SAIU = "Pronto! Você não vai mais receber nossas mensagens. Se mudar de ideia, responda VOLTAR. 🍦"
TEXTO_VOLTOU = "Que bom! 🍦 Você volta a receber as novidades do Gela Boca. Para sair, responda SAIR."
MODELOS = [
    ('🍦 Sabor novo', "Oi, {nome}! 🍦 Chegou sabor novo no Gela Boca: venha provar!"),
    ('🔥 Dia quente', "Oi, {nome}! 🔥 Hoje o calor está forte: que tal um sorvete bem gelado? Estamos te esperando!"),
    ('🌧️ Dia de chuva', "Oi, {nome}! 🌧️ Chovendo? A gente leva o sorvete até você: peça pelo delivery!"),
    ('🎉 Promoção', "Oi, {nome}! 🎉 Promoção especial no Gela Boca hoje. Mostre esta mensagem no caixa!"),
    ('🕐 Horário especial', "Oi, {nome}! 🕐 Aviso: no feriado o Gela Boca abre das __h às __h."),
]

_tabelas_ok = False
_trava_cadastro = threading.Lock()
_cadastros_por_ip = {}
_cadastros_do_dia = {'dia': None, 'n': 0}


class ErroClientes(Exception):
    def __init__(self, mensagem, status=400):
        super().__init__(mensagem)
        self.status = status


# ------------------------------------------------------------------------------
# Banco
# ------------------------------------------------------------------------------
def _conexao():
    conn = database.get_db_connection()
    if not conn:
        raise ErroClientes("Sem conexão com o banco de dados.", 503)
    return conn


def garantir_tabelas():
    global _tabelas_ok
    if _tabelas_ok:
        return
    conn = _conexao()
    try:
        cur = conn.cursor()
        cur.execute("""
            IF NOT EXISTS (SELECT * FROM sys.tables WHERE name = 'Clientes')
            CREATE TABLE Clientes (
                ClienteID INT IDENTITY(1,1) PRIMARY KEY,
                Nome NVARCHAR(120) NOT NULL,
                Telefone VARCHAR(20) NOT NULL UNIQUE,
                AniversarioDia INT NULL,
                AniversarioMes INT NULL,
                Aceite BIT NOT NULL DEFAULT 0,
                AceiteEm DATETIME NULL,
                SaiuEm DATETIME NULL,
                Origem VARCHAR(20) NULL,
                Observacao NVARCHAR(300) NULL,
                AniversarioEnviadoAno INT NULL,
                CriadoEm DATETIME NULL,
                AtualizadoEm DATETIME NULL
            )
        """)
        cur.execute("""
            IF NOT EXISTS (SELECT * FROM sys.tables WHERE name = 'ClienteEtiquetas')
            CREATE TABLE ClienteEtiquetas (
                ClienteID INT NOT NULL,
                Etiqueta NVARCHAR(40) NOT NULL,
                PRIMARY KEY (ClienteID, Etiqueta)
            )
        """)
        cur.execute("""
            IF NOT EXISTS (SELECT * FROM sys.tables WHERE name = 'CampanhasClientes')
            CREATE TABLE CampanhasClientes (
                CampanhaID INT IDENTITY(1,1) PRIMARY KEY,
                Tipo VARCHAR(20) NOT NULL,
                Titulo NVARCHAR(120) NULL,
                Texto NVARCHAR(2000) NOT NULL,
                Foto VARCHAR(60) NULL,
                Publico NVARCHAR(400) NULL,
                Status VARCHAR(20) NOT NULL,
                AgendadaPara DATETIME NULL,
                CriadaEm DATETIME NOT NULL,
                CriadaPor NVARCHAR(150) NULL,
                IniciadaEm DATETIME NULL,
                ConcluidaEm DATETIME NULL,
                Observacao NVARCHAR(300) NULL
            )
        """)
        cur.execute("""
            IF NOT EXISTS (SELECT * FROM sys.tables WHERE name = 'CampanhaEnvios')
            CREATE TABLE CampanhaEnvios (
                CampanhaID INT NOT NULL,
                ClienteID INT NOT NULL,
                Nome NVARCHAR(120) NULL,
                Telefone VARCHAR(20) NOT NULL,
                Status VARCHAR(12) NOT NULL,
                Erro NVARCHAR(200) NULL,
                EnviadoEm DATETIME NULL,
                PRIMARY KEY (CampanhaID, ClienteID)
            )
        """)
        cur.execute("""
            IF NOT EXISTS (SELECT * FROM sys.tables WHERE name = 'ParametrosClientes')
            CREATE TABLE ParametrosClientes (
                Chave VARCHAR(40) NOT NULL PRIMARY KEY,
                Valor NVARCHAR(400) NULL
            )
        """)
        conn.commit()
        _tabelas_ok = True
    finally:
        conn.close()


def _params(*chaves):
    garantir_tabelas()
    conn = _conexao()
    try:
        cur = conn.cursor()
        cur.execute("SELECT Chave, Valor FROM ParametrosClientes")
        todos = {c: v for c, v in cur.fetchall()}
    finally:
        conn.close()
    return {c: todos.get(c) for c in chaves} if chaves else todos


def _salvar_params(valores):
    garantir_tabelas()
    conn = _conexao()
    try:
        cur = conn.cursor()
        for chave, valor in valores.items():
            cur.execute("DELETE FROM ParametrosClientes WHERE Chave = ?", chave)
            if valor is not None:
                cur.execute("INSERT INTO ParametrosClientes (Chave, Valor) VALUES (?, ?)", chave, str(valor)[:400])
        conn.commit()
    finally:
        conn.close()


# ------------------------------------------------------------------------------
# Textos e números
# ------------------------------------------------------------------------------
def _limpo(texto, tamanho):
    return re.sub(r'\s+', ' ', str(texto if texto is not None else '')).strip()[:tamanho]


def _sem_acento(texto):
    return ''.join(c for c in unicodedata.normalize('NFD', str(texto or '')) if unicodedata.category(c) != 'Mn')


def _id(valor, nome="o cliente"):
    try:
        n = int(str(valor).strip())
    except (TypeError, ValueError):
        raise ErroClientes(f"Número inválido para {nome}. Atualize a página.")
    if n <= 0:
        raise ErroClientes(f"Número inválido para {nome}. Atualize a página.")
    return n


def normalizar_telefone(texto):
    """'(66) 99999-1234' / '66999991234' / '+55 66 9 9999-1234' -> '5566999991234'."""
    n = re.sub(r'\D', '', str(texto or '')).lstrip('0')
    if len(n) in (10, 11):
        n = '55' + n
    if not (n.startswith('55') and len(n) in (12, 13)):
        raise ErroClientes(f"WhatsApp inválido: '{_limpo(texto, 30)}'. Use o DDD + o número (ex.: 66 99999-1234).")
    return n


def fmt_telefone(n):
    """'5566999991234' -> '(66) 99999-1234'."""
    n = str(n or '')
    if n.startswith('55') and len(n) in (12, 13):
        ddd, resto = n[2:4], n[4:]
        return f"({ddd}) {resto[:-4]}-{resto[-4:]}"
    return n


def primeiro_nome(nome):
    p = _limpo(nome, 120).split(' ')[0] if _limpo(nome, 120) else ''
    return p[:1].upper() + p[1:].lower() if p else ''


def personalizar(texto, nome, rodape=True):
    """'Oi, {nome}!' -> 'Oi, Maria!' (sem nome: 'Oi!'), + o rodapé do SAIR."""
    pn = primeiro_nome(nome)
    t = re.sub(r'\{\s*nome\s*\}', pn, str(texto or ''), flags=re.I)
    if not pn:
        t = re.sub(r'[ \t]+([,!?.])', r'\1', t)
        t = re.sub(r',\s*([!?.])', r'\1', t)
    t = t.strip()
    return f"{t}\n\n{RODAPE}" if rodape else t


def _aniversario(valor):
    """'15/03' / '15/3' / date / datetime / '1990-03-15' -> (15, 3). Vazio -> (None, None)."""
    if valor in (None, ''):
        return None, None
    if isinstance(valor, (date, datetime)):
        return valor.day, valor.month
    t = str(valor).strip()
    m = re.fullmatch(r'(\d{1,2})\s*[/.-]\s*(\d{1,2})(?:\s*[/.-]\s*\d{2,4})?', t)
    if m:
        dia, mes = int(m.group(1)), int(m.group(2))
    else:
        m = re.fullmatch(r'\d{4}-(\d{2})-(\d{2})(?:[ T].*)?', t)
        if not m:
            raise ErroClientes(f"Aniversário inválido: '{_limpo(t, 20)}'. Use dia/mês (ex.: 15/03).")
        dia, mes = int(m.group(2)), int(m.group(1))
    try:
        date(2000, mes, dia)                       # 2000 é bissexto: 29/02 vale
    except ValueError:
        raise ErroClientes(f"Aniversário inválido: '{_limpo(t, 20)}'. Use dia/mês (ex.: 15/03).")
    return dia, mes


def _etiquetas(valor):
    """'VIP, Delivery' / ['VIP', 'vip '] -> ['VIP', 'Delivery'] (sem repetir, no máximo 10)."""
    itens = valor if isinstance(valor, list) else re.split(r'[,;]', str(valor or ''))
    saida, vistos = [], set()
    for e in itens:
        e = _limpo(e, 30)
        chave = _sem_acento(e).lower()
        if e and chave not in vistos:
            vistos.add(chave)
            saida.append(e)
    if len(saida) > MAX_ETIQUETAS:
        raise ErroClientes(f"No máximo {MAX_ETIQUETAS} etiquetas por cliente.")
    return saida


def _sem_segredo(texto):
    """Erros da biblioteca de internet trazem a URL com o token: some daqui antes de ir para log/tela."""
    t = str(texto)
    for v in _numero().values():
        if v and len(v) >= 6:
            t = t.replace(v, '***')
    return re.sub(r'/instances/[^/\s]+/token/[^/\s]+', '/instances/***/token/***', t)


# ------------------------------------------------------------------------------
# O número dos clientes (2ª instância da Z-API)
# ------------------------------------------------------------------------------
def _numero():
    p = _params('zapi_instancia', 'zapi_token', 'zapi_client_token')
    return {k: (v or '').strip() for k, v in p.items()}


def _base():
    n = _numero()
    if n['zapi_instancia'] and n['zapi_token']:
        return f"https://api.z-api.io/instances/{n['zapi_instancia']}/token/{n['zapi_token']}"
    return None


def configurado():
    try:
        return bool(_base())
    except ErroClientes:
        return False


def _cabecalhos():
    h = {"Content-Type": "application/json"}
    token = _numero()['zapi_client_token'] or getattr(config, 'ZAPI_CLIENT_TOKEN', '')   # mesma conta: mesmo Client-Token
    if token:
        h["Client-Token"] = token
    return h


def _requisicao(metodo, url, **kwargs):
    import requests
    return requests.request(metodo, url, **kwargs)


def _zapi(metodo, caminho, corpo=None, timeout=30):
    """(True, resposta) ou (False, motivo). Nunca usa o número da loja."""
    base = _base()
    if not base:
        return False, "o número dos clientes não está configurado (📢 Clientes › ⚙️ Ajustes)"
    try:
        r = _requisicao(metodo, f"{base}/{caminho}", json=corpo, headers=_cabecalhos(), timeout=timeout)
    except Exception as e:
        return False, _sem_segredo(f"sem conexão com a Z-API ({e})")[:180]
    try:
        dados = r.json()
    except ValueError:
        dados = {}
    if r.status_code != 200 or (isinstance(dados, dict) and dados.get('error')):
        motivo = (dados.get('error') or dados.get('message')) if isinstance(dados, dict) else ''
        return False, _sem_segredo(f"a Z-API recusou ({r.status_code}{': ' + str(motivo) if motivo else ''})")[:180]
    return True, dados


def salvar_numero(dados, usuario):
    """Grava o ID e o Token da 2ª instância (vazio no Token/Client-Token = mantém o que já estava)."""
    inst = _limpo(dados.get('instancia'), 80).replace(' ', '')
    token = _limpo(dados.get('token'), 80).replace(' ', '')
    client = _limpo(dados.get('client_token'), 80).replace(' ', '')
    if dados.get('apagar'):
        _salvar_params({'zapi_instancia': None, 'zapi_token': None, 'zapi_client_token': None})
        logger.info(f"[CLIENTES] {usuario.get('nome')}: apagou o número dos clientes.")
        return ajustes_numero()
    atual = _numero()
    if not inst:
        raise ErroClientes("Cole o ID da instância (Z-API › Instâncias › a do número dos clientes).")
    if not token and not atual['zapi_token']:
        raise ErroClientes("Cole o Token da instância.")
    if not re.fullmatch(r'[A-Za-z0-9_-]{6,80}', inst) or (token and not re.fullmatch(r'[A-Za-z0-9_-]{6,80}', token)):
        raise ErroClientes("O ID e o Token da Z-API só têm letras e números. Confira se colou certinho.")
    loja = (getattr(config, 'ZAPI_INSTANCE_ID', '') or '', getattr(config, 'WPP_API_URL', '') or '')
    if inst and (inst == loja[0] or f"/instances/{inst}/" in loja[1]):
        raise ErroClientes("Esse é o número da LOJA (o da escala e dos avisos). Use outra instância, com o chip só dos clientes.")
    novos = {'zapi_instancia': inst}
    if token:
        novos['zapi_token'] = token
    if client:
        novos['zapi_client_token'] = client
    _salvar_params(novos)
    logger.info(f"[CLIENTES] {usuario.get('nome')}: gravou o número dos clientes (instância …{inst[-4:]}).")
    return ajustes_numero()


def ajustes_numero():
    n = _numero()
    return {'configurado': bool(n['zapi_instancia'] and n['zapi_token']),
            'instancia_fim': n['zapi_instancia'][-4:] if n['zapi_instancia'] else '',
            'client_token_proprio': bool(n['zapi_client_token'])}


def testar_conexao():
    ok, r = _zapi('GET', 'status', timeout=15)
    if not ok:
        return {'ok': False, 'texto': f"❌ Não deu: {r}."}
    if isinstance(r, dict) and r.get('connected'):
        return {'ok': True, 'texto': "✅ Número dos clientes conectado e pronto para enviar."}
    return {'ok': False, 'texto': "⚠️ A instância existe, mas o celular não está conectado. No painel da Z-API, "
                                  "leia o QR Code com o WhatsApp do chip dos clientes."}


# ------------------------------------------------------------------------------
# Enviar UMA mensagem (texto ou foto + legenda)
# ------------------------------------------------------------------------------
def caminho_foto(nome):
    if not re.fullmatch(r'[0-9a-f]{32}\.(jpg|png|webp)', str(nome or '')):
        raise ErroClientes("Foto não encontrada.", 404)
    caminho = os.path.join(PASTA_FOTOS, nome)
    if not os.path.isfile(caminho):
        raise ErroClientes("Foto não encontrada (foi apagada?). Escolha a foto de novo.", 404)
    return caminho


def salvar_foto(data_url):
    """'data:image/jpeg;base64,...' (a página já diminui a foto) -> nome do arquivo guardado."""
    m = re.fullmatch(r'data:image/(jpeg|jpg|png|webp);base64,([A-Za-z0-9+/=\s]+)', str(data_url or '').strip())
    if not m:
        raise ErroClientes("Foto inválida. Tire ou escolha a foto de novo.")
    try:
        dados = base64.b64decode(m.group(2), validate=False)
    except (ValueError, TypeError):
        raise ErroClientes("Foto inválida. Tire ou escolha a foto de novo.")
    if len(dados) > MAX_FOTO_BYTES:
        raise ErroClientes("Foto grande demais (máximo 3 MB).")
    if dados[:3] == b'\xff\xd8\xff':
        ext = 'jpg'
    elif dados[:8] == b'\x89PNG\r\n\x1a\n':
        ext = 'png'
    elif dados[:4] == b'RIFF' and dados[8:12] == b'WEBP':
        ext = 'webp'
    else:
        raise ErroClientes("O arquivo não é uma foto (JPG, PNG ou WEBP).")
    os.makedirs(PASTA_FOTOS, exist_ok=True)
    nome = f"{uuid.uuid4().hex}.{ext}"
    with open(os.path.join(PASTA_FOTOS, nome), 'wb') as f:
        f.write(dados)
    return {'foto': nome}


def enviar_para(telefone, texto, foto=None):
    """(True, '') ou (False, motivo). Foto: send-image com legenda (texto longo vai em seguida, separado)."""
    if not foto:
        ok, r = _zapi('POST', 'send-text', {"phone": telefone, "message": texto})
        return (True, '') if ok else (False, r)
    caminho = caminho_foto(foto)
    with open(caminho, 'rb') as f:
        b64 = base64.b64encode(f.read()).decode()
    mime = {'jpg': 'jpeg', 'png': 'png', 'webp': 'webp'}[foto.rsplit('.', 1)[1]]
    legenda = texto if len(texto) <= MAX_LEGENDA else ''
    ok, r = _zapi('POST', 'send-image', {"phone": telefone, "image": f"data:image/{mime};base64,{b64}", "caption": legenda}, timeout=60)
    if not ok:
        return False, r
    if not legenda:
        ok, r = _zapi('POST', 'send-text', {"phone": telefone, "message": texto})
        if not ok:
            return False, f"a foto foi, o texto não: {r}"
    return True, ''


def _em_segundo_plano(funcao, *args):
    threading.Thread(target=funcao, args=args, daemon=True).start()


# ------------------------------------------------------------------------------
# Cadastro dos clientes
# ------------------------------------------------------------------------------
def _recebe(row_aceite, row_saiu):
    return bool(row_aceite) and not row_saiu


def _todas_etiquetas(cur):
    cur.execute("SELECT ClienteID, Etiqueta FROM ClienteEtiquetas")
    por = {}
    for cid, e in cur.fetchall():
        por.setdefault(cid, []).append(e)
    return por


COLUNAS = "ClienteID, Nome, Telefone, AniversarioDia, AniversarioMes, Aceite, AceiteEm, SaiuEm, Origem, Observacao, CriadoEm"


def _cliente_json(r, etqs):
    cid, nome, tel, dia, mes, aceite, aceite_em, saiu, origem, obs, criado = r
    situacao = 'saiu' if saiu else ('recebe' if aceite else 'sem_autorizacao')
    return {'id': cid, 'nome': nome, 'telefone': tel, 'telefone_fmt': fmt_telefone(tel),
            'aniversario': f"{dia:02d}/{mes:02d}" if dia and mes else '', 'etiquetas': sorted(etqs, key=str.lower),
            'situacao': situacao, 'origem': origem or '', 'observacao': obs or '',
            'saiu_em': database._como_data(saiu).strftime('%d/%m/%Y') if saiu else ''}


def listar(filtros=None):
    garantir_tabelas()
    filtros = filtros or {}
    busca = _sem_acento(_limpo(filtros.get('busca'), 60)).lower()
    etiqueta = _limpo(filtros.get('etiqueta'), 30)
    situacao = filtros.get('situacao') if filtros.get('situacao') in ('recebe', 'saiu', 'sem_autorizacao') else 'todos'
    conn = _conexao()
    try:
        cur = conn.cursor()
        cur.execute(f"SELECT {COLUNAS} FROM Clientes")
        linhas = cur.fetchall()
        etqs = _todas_etiquetas(cur)
    finally:
        conn.close()
    todos = [_cliente_json(r, etqs.get(r[0], [])) for r in linhas]
    contagem = {}
    for c in todos:
        for e in c['etiquetas']:
            contagem.setdefault(e, [0, 0])
            contagem[e][0] += 1
            contagem[e][1] += 1 if c['situacao'] == 'recebe' else 0
    sel = todos
    if busca:
        dig = re.sub(r'\D', '', busca)
        sel = [c for c in sel if busca in _sem_acento(c['nome']).lower() or (len(dig) >= 3 and dig in c['telefone'])]
    if etiqueta:
        sel = [c for c in sel if any(_sem_acento(e).lower() == _sem_acento(etiqueta).lower() for e in c['etiquetas'])]
    if situacao != 'todos':
        sel = [c for c in sel if c['situacao'] == situacao]
    sel.sort(key=lambda c: _sem_acento(c['nome']).lower())
    return {'clientes': sel[:1000], 'total_filtro': len(sel),
            'totais': {'todos': len(todos), 'recebe': sum(1 for c in todos if c['situacao'] == 'recebe'),
                       'saiu': sum(1 for c in todos if c['situacao'] == 'saiu'),
                       'sem_autorizacao': sum(1 for c in todos if c['situacao'] == 'sem_autorizacao')},
            'etiquetas': [{'nome': e, 'qtd': q, 'recebem': r} for e, (q, r) in sorted(contagem.items(), key=lambda kv: kv[0].lower())],
            'configurado': configurado()}


def _gravar_etiquetas(cur, cid, etiquetas, somar=False):
    if not somar:
        cur.execute("DELETE FROM ClienteEtiquetas WHERE ClienteID = ?", cid)
        atuais = set()
    else:
        cur.execute("SELECT Etiqueta FROM ClienteEtiquetas WHERE ClienteID = ?", cid)
        atuais = {_sem_acento(r[0]).lower() for r in cur.fetchall()}
    for e in etiquetas:
        if _sem_acento(e).lower() not in atuais:
            cur.execute("INSERT INTO ClienteEtiquetas (ClienteID, Etiqueta) VALUES (?, ?)", cid, e)
            atuais.add(_sem_acento(e).lower())


def _por_telefone(cur, telefone):
    cur.execute("SELECT ClienteID, Nome, Aceite, SaiuEm FROM Clientes WHERE Telefone = ?", telefone)
    return cur.fetchone()


def salvar(cliente_id, dados, usuario):
    """Novo (cliente_id vazio) ou editar. 'aceite' = o cliente autorizou receber mensagens."""
    garantir_tabelas()
    nome = _limpo(dados.get('nome'), 120)
    if len(nome) < 2:
        raise ErroClientes("Digite o nome do cliente.")
    tel = normalizar_telefone(dados.get('telefone'))
    dia, mes = _aniversario(dados.get('aniversario'))
    etiquetas = _etiquetas(dados.get('etiquetas'))
    aceite = bool(dados.get('aceite'))
    obs = _limpo(dados.get('observacao'), 300) or None
    agora = datetime.now().replace(microsecond=0)
    cid = _id(cliente_id) if cliente_id not in (None, '', 0, '0') else None
    conn = _conexao()
    try:
        cur = conn.cursor()
        outro = _por_telefone(cur, tel)
        if outro and outro[0] != cid:
            raise ErroClientes(f"Esse WhatsApp já está cadastrado para {outro[1]}.", 409)
        if cid:
            cur.execute("SELECT Aceite FROM Clientes WHERE ClienteID = ?", cid)
            antes = cur.fetchone()
            if not antes:
                raise ErroClientes("Cliente não encontrado (foi excluído?). Atualize a página.", 409)
            cur.execute("UPDATE Clientes SET Nome = ?, Telefone = ?, AniversarioDia = ?, AniversarioMes = ?, Aceite = ?, "
                        "AceiteEm = CASE WHEN ? = 1 AND Aceite = 0 THEN ? ELSE AceiteEm END, Observacao = ?, AtualizadoEm = ? "
                        "WHERE ClienteID = ?", nome, tel, dia, mes, 1 if aceite else 0, 1 if aceite else 0, agora, obs, agora, cid)
        else:
            cur.execute("INSERT INTO Clientes (Nome, Telefone, AniversarioDia, AniversarioMes, Aceite, AceiteEm, Origem, Observacao, "
                        "CriadoEm, AtualizadoEm) VALUES (?, ?, ?, ?, ?, ?, 'manual', ?, ?, ?); SELECT SCOPE_IDENTITY();",
                        nome, tel, dia, mes, 1 if aceite else 0, agora if aceite else None, obs, agora, agora)
            cur.nextset()
            cid = int(cur.fetchone()[0])
        _gravar_etiquetas(cur, cid, etiquetas)
        conn.commit()
    finally:
        conn.close()
    logger.info(f"[CLIENTES] {usuario.get('nome')}: salvou o cliente {cid} ({nome}).")
    return {'ok': True, 'id': cid, 'mensagem': f"{nome} salvo." + ('' if aceite else ' Sem autorização: não recebe as mensagens.')}


def excluir(cliente_id, usuario):
    cid = _id(cliente_id)
    garantir_tabelas()
    conn = _conexao()
    try:
        cur = conn.cursor()
        cur.execute("SELECT Nome FROM Clientes WHERE ClienteID = ?", cid)
        r = cur.fetchone()
        if not r:
            raise ErroClientes("Cliente não encontrado (já foi excluído?). Atualize a página.", 409)
        cur.execute("DELETE FROM ClienteEtiquetas WHERE ClienteID = ?", cid)
        cur.execute("UPDATE CampanhaEnvios SET Status = 'pulado', Erro = 'cliente excluído' WHERE ClienteID = ? AND Status = 'pendente'", cid)
        cur.execute("DELETE FROM Clientes WHERE ClienteID = ?", cid)
        conn.commit()
    finally:
        conn.close()
    logger.info(f"[CLIENTES] {usuario.get('nome')}: excluiu o cliente {cid} ({r[0]}).")
    return {'ok': True, 'mensagem': f"{r[0]} excluído."}


def definir_saida(cliente_id, saiu, usuario):
    """O cliente pediu para sair (não recebe mais) ou voltou."""
    cid = _id(cliente_id)
    garantir_tabelas()
    agora = datetime.now().replace(microsecond=0)
    conn = _conexao()
    try:
        cur = conn.cursor()
        cur.execute("SELECT Nome, Aceite FROM Clientes WHERE ClienteID = ?", cid)
        r = cur.fetchone()
        if not r:
            raise ErroClientes("Cliente não encontrado. Atualize a página.", 409)
        if saiu:
            cur.execute("UPDATE Clientes SET SaiuEm = ?, AtualizadoEm = ? WHERE ClienteID = ?", agora, agora, cid)
        else:
            cur.execute("UPDATE Clientes SET SaiuEm = NULL, AtualizadoEm = ? WHERE ClienteID = ?", agora, cid)
        conn.commit()
    finally:
        conn.close()
    logger.info(f"[CLIENTES] {usuario.get('nome')}: {r[0]} {'saiu da lista' if saiu else 'voltou para a lista'}.")
    return {'ok': True, 'mensagem': f"{r[0]}: {'não recebe mais as mensagens' if saiu else 'volta a receber'}"
                                    + ('' if saiu or r[1] else ' (falta marcar que autorizou)') + '.'}


# ------------------------------------------------------------------------------
# Mensagens (campanhas): para quem, prévia, criar, acompanhar
# ------------------------------------------------------------------------------
def _publico(dados):
    p = dados.get('publico') or {}
    if not isinstance(p, dict):
        raise ErroClientes("Escolha para quem vai a mensagem.")
    etiquetas = _etiquetas(p.get('etiquetas') or [])
    todos = bool(p.get('todos'))
    if not todos and not etiquetas:
        raise ErroClientes("Escolha para quem vai a mensagem: todos ou pelo menos uma etiqueta.")
    return {'todos': todos, 'etiquetas': [] if todos else etiquetas}


def _texto_publico(p):
    return "Todos os clientes" if p['todos'] else "Etiquetas: " + ", ".join(p['etiquetas'])


def destinatarios(publico, cur=None):
    """[(ClienteID, Nome, Telefone)] de quem RECEBE (autorizou e não saiu) e está no público."""
    garantir_tabelas()
    proprio = cur is None
    conn = _conexao() if proprio else None
    try:
        cur = cur or conn.cursor()
        cur.execute("SELECT ClienteID, Nome, Telefone, Aceite, SaiuEm FROM Clientes")
        clientes = [r for r in cur.fetchall() if _recebe(r[3], r[4])]
        etqs = _todas_etiquetas(cur) if not publico['todos'] else {}
    finally:
        if proprio:
            conn.close()
    if not publico['todos']:
        alvo = {_sem_acento(e).lower() for e in publico['etiquetas']}
        clientes = [r for r in clientes if alvo & {_sem_acento(e).lower() for e in etqs.get(r[0], [])}]
    return sorted([(r[0], r[1], r[2]) for r in clientes], key=lambda r: _sem_acento(r[1]).lower())


def _texto(dados):
    t = str(dados.get('texto') or '').replace('\r\n', '\n').strip()
    if not t:
        raise ErroClientes("Escreva a mensagem.")
    if len(t) > MAX_TEXTO:
        raise ErroClientes(f"Mensagem longa demais ({len(t)} letras; máximo {MAX_TEXTO}).")
    return t


def _foto_dados(dados):
    foto = dados.get('foto') or None
    if foto:
        caminho_foto(foto)
    return foto


def _minutos(qtd):
    return max(1, round(qtd * sum(INTERVALO_SEG) / 2 / 60))


def enviados_hoje(cur, agora):
    ini = datetime.combine(agora.date(), datetime.min.time())
    cur.execute("SELECT COUNT(*) FROM CampanhaEnvios WHERE Status = 'enviado' AND EnviadoEm >= ?", ini)
    return int(cur.fetchone()[0] or 0)


def previa(dados):
    """Quantos vão receber, como fica a mensagem (com o 1º nome) e quanto tempo leva."""
    p = _publico(dados)
    texto = str(dados.get('texto') or '').strip()
    garantir_tabelas()
    conn = _conexao()
    try:
        cur = conn.cursor()
        lista = destinatarios(p, cur)
        hoje = enviados_hoje(cur, datetime.now())
    finally:
        conn.close()
    exemplo = lista[0][1] if lista else 'Maria'
    return {'qtd': len(lista), 'exemplo_nome': primeiro_nome(exemplo), 'texto': personalizar(texto, exemplo) if texto else '',
            'minutos': _minutos(len(lista)) if lista else 0, 'restante_hoje': max(0, LIMITE_POR_DIA - hoje),
            'nomes': [primeiro_nome(r[1]) or fmt_telefone(r[2]) for r in lista[:6]], 'configurado': configurado()}


def _quando(dados, agora):
    txt = str(dados.get('agendar') or '').strip()
    if not txt:
        return None
    try:
        quando = datetime.fromisoformat(txt[:16])
    except ValueError:
        raise ErroClientes("Data e hora do agendamento inválidas.")
    if quando <= agora + timedelta(minutes=1):
        raise ErroClientes("Escolha um horário no futuro para agendar (ou envie agora).")
    if quando > agora + timedelta(days=60):
        raise ErroClientes("Agende para no máximo 60 dias à frente.")
    if not (HORARIO_ENVIO[0] <= quando.hour < HORARIO_ENVIO[1]):
        raise ErroClientes(f"As mensagens só saem das {HORARIO_ENVIO[0]}h às {HORARIO_ENVIO[1]}h (para não incomodar). "
                           "Escolha um horário nessa faixa.")
    return quando


def _criar(cur, tipo, titulo, texto, foto, publico_txt, lista, status, agendada, usuario_nome, agora):
    cur.execute("INSERT INTO CampanhasClientes (Tipo, Titulo, Texto, Foto, Publico, Status, AgendadaPara, CriadaEm, CriadaPor) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?); SELECT SCOPE_IDENTITY();",
                tipo, titulo[:120], texto, foto, publico_txt[:400], status, agendada, agora, usuario_nome)
    cur.nextset()
    camp = int(cur.fetchone()[0])
    for cid, nome, tel in lista:
        cur.execute("INSERT INTO CampanhaEnvios (CampanhaID, ClienteID, Nome, Telefone, Status) VALUES (?, ?, ?, ?, 'pendente')",
                    camp, cid, nome, tel)
    return camp


def criar_campanha(dados, usuario):
    """Grava a mensagem e a lista de quem recebe; o robô manda (agora ou na hora agendada)."""
    texto, foto, p = _texto(dados), _foto_dados(dados), _publico(dados)
    agora = datetime.now().replace(second=0, microsecond=0)
    quando = _quando(dados, agora)
    if not configurado():
        raise ErroClientes("Antes, configure o número dos clientes em ⚙️ Ajustes (a 2ª instância da Z-API).")
    confirmados = set(dados.get('confirmar') or [])
    garantir_tabelas()
    conn = _conexao()
    try:
        cur = conn.cursor()
        lista = destinatarios(p, cur)
        if not lista:
            raise ErroClientes("Ninguém para receber: nenhum cliente dessas etiquetas autorizou receber mensagens.")
        perguntas = []
        cur.execute("SELECT CampanhaID, CriadaEm, Texto FROM CampanhasClientes WHERE CriadaEm >= ?", agora - timedelta(minutes=10))
        if 'repetida' not in confirmados and any(r[2] == texto for r in cur.fetchall()):
            perguntas.append({'chave': 'repetida', 'titulo': "Mensagem repetida",
                              'texto': "Você criou esta mesma mensagem há poucos minutos.\n\nMandar de novo mesmo assim?"})
        restante = max(0, LIMITE_POR_DIA - enviados_hoje(cur, agora))
        if quando is None and len(lista) > restante and 'limite' not in confirmados:
            perguntas.append({'chave': 'limite', 'titulo': "Limite do dia", 'texto': (
                f"Hoje ainda dá para mandar {restante} mensagem(ns) (limite de {LIMITE_POR_DIA} por dia, para o WhatsApp "
                f"não bloquear o número). As outras {len(lista) - restante} saem amanhã a partir das {HORARIO_ENVIO[0]}h.\n\nContinuar?")})
        if perguntas:
            return {'perguntas': perguntas}
        titulo = _limpo(re.sub(r'\{\s*nome\s*\}', '…', texto.split('\n')[0], flags=re.I), 80)
        camp = _criar(cur, 'mensagem', titulo, texto, foto, _texto_publico(p), lista,
                      ST_AGENDADA if quando else ST_ENVIANDO, quando, usuario.get('nome'), agora)
        conn.commit()
    finally:
        conn.close()
    logger.info(f"[CLIENTES] {usuario.get('nome')}: mensagem {camp} para {len(lista)} cliente(s)"
                + (f", agendada para {quando:%d/%m %H:%M}" if quando else '') + '.')
    if quando:
        msg = f"Agendada para {quando:%d/%m às %H:%M}: {len(lista)} cliente(s)."
    else:
        msg = (f"Na fila: {len(lista)} cliente(s). O robô começa em até 1 minuto e leva cerca de {_minutos(len(lista))} min "
               f"(uma mensagem a cada {INTERVALO_SEG[0]} a {INTERVALO_SEG[1]} s).")
    return {'ok': True, 'id': camp, 'mensagem': msg}


def enviar_teste(dados):
    """Manda a mensagem AGORA para um número (o seu), para ver como chega. Não entra no histórico."""
    texto, foto = _texto(dados), _foto_dados(dados)
    tel = normalizar_telefone(dados.get('telefone'))
    if not configurado():
        raise ErroClientes("Antes, configure o número dos clientes em ⚙️ Ajustes (a 2ª instância da Z-API).")
    ok, motivo = enviar_para(tel, personalizar(texto, dados.get('nome') or 'Rodrigo'), foto)
    if not ok:
        raise ErroClientes(f"O teste não foi: {motivo}.")
    return {'ok': True, 'mensagem': f"Teste enviado para {fmt_telefone(tel)}. Confira no WhatsApp."}


def _contagens(cur):
    cur.execute("SELECT CampanhaID, Status, COUNT(*) FROM CampanhaEnvios GROUP BY CampanhaID, Status")
    por = {}
    for camp, st, n in cur.fetchall():
        por.setdefault(camp, {})[st] = n
    return por


def _campanha_json(r, cont):
    camp, tipo, titulo, texto, foto, publico, status, agendada, criada, por, iniciada, concluida, obs = r
    c = cont.get(camp, {})
    total = sum(c.values())
    return {'id': camp, 'tipo': tipo, 'titulo': titulo or '', 'texto': texto, 'foto': foto or '', 'publico': publico or '',
            'status': status, 'status_nome': NOMES_STATUS.get(status, status),
            'agendada': _dt_txt(agendada), 'criada': _dt_txt(criada), 'criada_por': por or '',
            'concluida': _dt_txt(concluida) if concluida else '', 'observacao': obs or '',
            'total': total, 'enviados': c.get('enviado', 0), 'falhas': c.get('falhou', 0), 'pulados': c.get('pulado', 0),
            'pendentes': c.get('pendente', 0)}


def _dt(v):
    if v is None or isinstance(v, datetime):
        return v
    try:
        return datetime.fromisoformat(str(v)[:19])
    except ValueError:
        return None


def _dt_txt(v):
    v = _dt(v)
    return v.strftime('%d/%m %H:%M') if v else ''


CAMPOS_CAMPANHA = ("CampanhaID, Tipo, Titulo, Texto, Foto, Publico, Status, AgendadaPara, CriadaEm, CriadaPor, IniciadaEm, "
                   "ConcluidaEm, Observacao")


def campanhas():
    garantir_tabelas()
    conn = _conexao()
    try:
        cur = conn.cursor()
        cur.execute(f"SELECT {CAMPOS_CAMPANHA} FROM CampanhasClientes")
        linhas = cur.fetchall()
        cont = _contagens(cur)
        hoje = enviados_hoje(cur, datetime.now())
    finally:
        conn.close()
    lista = sorted((_campanha_json(r, cont) for r in linhas), key=lambda c: -c['id'])[:60]
    return {'campanhas': lista, 'enviados_hoje': hoje, 'limite_dia': LIMITE_POR_DIA, 'configurado': configurado(),
            'horario': f"{HORARIO_ENVIO[0]}h às {HORARIO_ENVIO[1]}h"}


def campanha(camp_id):
    camp = _id(camp_id, "a mensagem")
    garantir_tabelas()
    conn = _conexao()
    try:
        cur = conn.cursor()
        cur.execute(f"SELECT {CAMPOS_CAMPANHA} FROM CampanhasClientes WHERE CampanhaID = ?", camp)
        r = cur.fetchone()
        if not r:
            raise ErroClientes("Mensagem não encontrada. Atualize a página.", 404)
        cont = _contagens(cur)
        cur.execute("SELECT ClienteID, Nome, Telefone, Status, Erro, EnviadoEm FROM CampanhaEnvios WHERE CampanhaID = ?", camp)
        envios = [{'cliente_id': e[0], 'nome': e[1] or '', 'telefone_fmt': fmt_telefone(e[2]), 'status': e[3],
                   'erro': e[4] or '', 'enviado': _dt_txt(e[5])} for e in cur.fetchall()]
    finally:
        conn.close()
    ordem = {'falhou': 0, 'pendente': 1, 'pulado': 2, 'enviado': 3}
    envios.sort(key=lambda e: (ordem.get(e['status'], 9), _sem_acento(e['nome']).lower()))
    return dict(_campanha_json(r, cont), envios=envios)


def acao_campanha(camp_id, acao, usuario):
    """cancelar (agendada/enviando/pausada) · continuar (pausada) · reenviar (quem falhou volta para a fila)."""
    camp = _id(camp_id, "a mensagem")
    if acao not in ('cancelar', 'continuar', 'reenviar'):
        raise ErroClientes("Ação desconhecida.")
    garantir_tabelas()
    agora = datetime.now().replace(microsecond=0)
    conn = _conexao()
    try:
        cur = conn.cursor()
        cur.execute("SELECT Status FROM CampanhasClientes WHERE CampanhaID = ?", camp)
        r = cur.fetchone()
        if not r:
            raise ErroClientes("Mensagem não encontrada. Atualize a página.", 404)
        st = r[0]
        if acao == 'cancelar':
            if st not in (ST_AGENDADA, ST_ENVIANDO, ST_PAUSADA):
                raise ErroClientes("Essa mensagem já terminou.", 409)
            cur.execute("UPDATE CampanhaEnvios SET Status = 'pulado', Erro = 'cancelada' WHERE CampanhaID = ? AND Status = 'pendente'", camp)
            cur.execute("UPDATE CampanhasClientes SET Status = ?, ConcluidaEm = ?, Observacao = ? WHERE CampanhaID = ?",
                        ST_CANCELADA, agora, f"Cancelada por {usuario.get('nome')}", camp)
            msg = "Cancelada: quem ainda não recebeu não vai receber."
        elif acao == 'continuar':
            if st != ST_PAUSADA:
                raise ErroClientes("Só dá para continuar uma mensagem pausada.", 409)
            cur.execute("UPDATE CampanhasClientes SET Status = ?, Observacao = NULL WHERE CampanhaID = ?", ST_ENVIANDO, camp)
            msg = "Continuando: o robô volta a mandar em até 1 minuto."
        else:
            if st == ST_CANCELADA:
                raise ErroClientes("Essa mensagem foi cancelada.", 409)
            cur.execute("SELECT COUNT(*) FROM CampanhaEnvios WHERE CampanhaID = ? AND Status = 'falhou'", camp)
            n = int(cur.fetchone()[0] or 0)
            if not n:
                raise ErroClientes("Ninguém falhou nessa mensagem.", 409)
            cur.execute("UPDATE CampanhaEnvios SET Status = 'pendente', Erro = NULL WHERE CampanhaID = ? AND Status = 'falhou'", camp)
            cur.execute("UPDATE CampanhasClientes SET Status = ?, ConcluidaEm = NULL, Observacao = NULL WHERE CampanhaID = ?",
                        ST_ENVIANDO, camp)
            msg = f"{n} cliente(s) de volta na fila."
        conn.commit()
    finally:
        conn.close()
    logger.info(f"[CLIENTES] {usuario.get('nome')}: {acao} a mensagem {camp}.")
    return {'ok': True, 'mensagem': msg}


# ------------------------------------------------------------------------------
# O robô: manda a fila (agendador.py, de minuto em minuto)
# ------------------------------------------------------------------------------
def _sai_ainda(cur, cid):
    cur.execute("SELECT Aceite, SaiuEm FROM Clientes WHERE ClienteID = ?", cid)
    r = cur.fetchone()
    return bool(r) and _recebe(r[0], r[1])


def processar_fila(agora=None, dormir=time.sleep, avisar=None):
    """
    Manda as mensagens que chegaram na hora, uma de cada vez, com intervalo aleatório; respeita o horário
    (8h–21h), o limite do dia e para sozinho se falhar 5 vezes seguidas. Devolve quantas mandou.
    'agora' (testes) = função que devolve o datetime atual.
    """
    relogio = agora or datetime.now
    avisar = avisar or (lambda texto: None)
    if not configurado():
        return 0
    garantir_tabelas()
    mandadas = 0
    while True:
        momento = relogio()
        if not (HORARIO_ENVIO[0] <= momento.hour < HORARIO_ENVIO[1]):
            return mandadas
        conn = _conexao()
        try:
            cur = conn.cursor()
            cur.execute("SELECT CampanhaID, Status, AgendadaPara, Texto, Foto, Titulo, Tipo FROM CampanhasClientes "
                        "WHERE Status IN (?, ?)", ST_AGENDADA, ST_ENVIANDO)
            prontas = sorted([r for r in cur.fetchall() if r[1] == ST_ENVIANDO or (_dt(r[2]) and _dt(r[2]) <= momento)],
                             key=lambda r: (_dt(r[2]) or datetime.min, r[0]))
            if not prontas:
                return mandadas
            camp, st, _ag, texto, foto, titulo, tipo = prontas[0]
            if st == ST_AGENDADA:
                cur.execute("UPDATE CampanhasClientes SET Status = ?, IniciadaEm = ? WHERE CampanhaID = ?", ST_ENVIANDO, momento, camp)
                conn.commit()
            cur.execute("SELECT ClienteID, Nome, Telefone FROM CampanhaEnvios WHERE CampanhaID = ? AND Status = 'pendente'", camp)
            fila = sorted(cur.fetchall(), key=lambda r: (_sem_acento(r[1] or '').lower(), r[0]))
            if not fila:
                cur.execute("UPDATE CampanhasClientes SET Status = ?, ConcluidaEm = ? WHERE CampanhaID = ?", ST_CONCLUIDA, momento, camp)
                conn.commit()
                cont = _contagens(cur).get(camp, {})
                cur.execute("SELECT Nome FROM CampanhaEnvios WHERE CampanhaID = ? AND Status = 'enviado'", camp)
                nomes = sorted((primeiro_nome(r[0]) for r in cur.fetchall() if r[0]), key=lambda n: _sem_acento(n).lower())
                avisar(_texto_resumo(titulo, cont, nomes))
                continue
            if enviados_hoje(cur, momento) >= LIMITE_POR_DIA:
                _avisar_limite(momento, avisar)
                return mandadas
        finally:
            conn.close()
        falhas_seguidas = 0
        for i, (cid, nome, tel) in enumerate(fila):
            conn = _conexao()
            try:
                cur = conn.cursor()
                cur.execute("SELECT Status FROM CampanhasClientes WHERE CampanhaID = ?", camp)
                if (cur.fetchone() or [None])[0] != ST_ENVIANDO:
                    break                                    # cancelaram/pausaram pela tela
                momento = relogio()
                if not (HORARIO_ENVIO[0] <= momento.hour < HORARIO_ENVIO[1]):
                    break
                if enviados_hoje(cur, momento) >= LIMITE_POR_DIA:
                    _avisar_limite(momento, avisar)
                    break
                cur.execute("SELECT Status FROM CampanhaEnvios WHERE CampanhaID = ? AND ClienteID = ?", camp, cid)
                if (cur.fetchone() or [None])[0] != 'pendente':
                    continue
                if not _sai_ainda(cur, cid):
                    cur.execute("UPDATE CampanhaEnvios SET Status = 'pulado', Erro = 'pediu para sair' WHERE CampanhaID = ? AND ClienteID = ?",
                                camp, cid)
                    conn.commit()
                    continue
            finally:
                conn.close()
            try:
                ok, motivo = enviar_para(tel, personalizar(texto, nome), foto)
            except ErroClientes as e:                        # a foto sumiu da pasta
                ok, motivo = False, str(e)
            except Exception as e:
                ok, motivo = False, _sem_segredo(e)[:180]
            conn = _conexao()
            try:
                cur = conn.cursor()
                cur.execute("UPDATE CampanhaEnvios SET Status = ?, Erro = ?, EnviadoEm = ? WHERE CampanhaID = ? AND ClienteID = ?",
                            'enviado' if ok else 'falhou', None if ok else str(motivo)[:200], relogio().replace(microsecond=0), camp, cid)
                falhas_seguidas = 0 if ok else falhas_seguidas + 1
                if falhas_seguidas >= FALHAS_SEGUIDAS:
                    cur.execute("UPDATE CampanhasClientes SET Status = ?, Observacao = ? WHERE CampanhaID = ?", ST_PAUSADA,
                                f"Pausada: {FALHAS_SEGUIDAS} falhas seguidas ({str(motivo)[:120]})", camp)
                conn.commit()
            finally:
                conn.close()
            if ok:
                mandadas += 1
            if falhas_seguidas >= FALHAS_SEGUIDAS:
                logger.error(f"[CLIENTES] Mensagem {camp} pausada: {FALHAS_SEGUIDAS} falhas seguidas ({motivo}).")
                avisar(f"⚠️ <b>Envio para clientes PAUSADO</b> (número dos clientes{_fim_instancia()}): «{_esc(titulo)}»\n"
                       f"{FALHAS_SEGUIDAS} falhas seguidas ({_esc(motivo)}). O número dos clientes pode ter desconectado: "
                       "confira na Z-API e toque em ▶️ Continuar em Gestão › 📢 Clientes › Envios.")
                return mandadas
            if i < len(fila) - 1:
                dormir(random.uniform(*INTERVALO_SEG))
        else:
            continue                                         # terminou a fila desta: volta e marca concluída
        return mandadas


def _fim_instancia():
    inst = _numero()['zapi_instancia']
    return f", instância …{inst[-4:]}" if inst else ''


def _texto_resumo(titulo, cont, nomes):
    """
    [CLIENTES] Resumo para o GRUPO DA GESTÃO (vai pelo canal dos avisos, como o boletim). Deixa claro que a mensagem
    foi direto para cada cliente pelo número dos clientes, e não para o grupo.
    """
    n = cont.get('enviado', 0)
    quem = ''
    if nomes:
        quem = ': ' + ', '.join(_esc(x) for x in nomes[:8]) + (f" e mais {len(nomes) - 8}" if len(nomes) > 8 else '')
    return (f"📢 <b>Resumo do envio para clientes</b> (só para a gestão)\n"
            f"«{_esc(titulo)}»\n"
            f"✅ {n} cliente(s) receberam no WhatsApp deles{quem}"
            + (f"\n❌ {cont['falhou']} falharam" if cont.get('falhou') else '')
            + (f"\n⏭️ {cont['pulado']} pularam (saíram da lista)" if cont.get('pulado') else '')
            + f"\n📱 Saiu pelo número dos clientes{_fim_instancia()}, direto para cada um."
            + "\n<i>Detalhes em Gestão › 📢 Clientes › Envios. Para não receber este resumo: 📣 Avisos.</i>")


def _avisar_limite(momento, avisar):
    """Chegou no limite do dia: avisa a gestão uma vez só por dia."""
    if _params('limite_avisado').get('limite_avisado') != momento.date().isoformat():
        _salvar_params({'limite_avisado': momento.date().isoformat()})
        avisar(f"📢 Limite de {LIMITE_POR_DIA} mensagens para clientes hoje: o resto sai amanhã a partir das "
               f"{HORARIO_ENVIO[0]}h (proteção contra bloqueio do número).")


def _esc(t):
    import html
    return html.escape(str(t or ''), quote=False)


# ------------------------------------------------------------------------------
# 🎂 Aniversário automático (agendador.py, 10h)
# ------------------------------------------------------------------------------
def config_aniversario():
    p = _params('aniversario_ligado', 'aniversario_texto', 'aniversario_foto')
    return {'ligado': p.get('aniversario_ligado') == '1', 'texto': p.get('aniversario_texto') or TEXTO_ANIVERSARIO_PADRAO,
            'foto': p.get('aniversario_foto') or ''}


def salvar_config_aniversario(dados, usuario):
    texto = _texto(dados)
    foto = _foto_dados(dados)
    _salvar_params({'aniversario_ligado': '1' if dados.get('ligado') else '0', 'aniversario_texto': texto[:400],
                    'aniversario_foto': foto})
    logger.info(f"[CLIENTES] {usuario.get('nome')}: aniversário {'ligado' if dados.get('ligado') else 'desligado'}.")
    return dict(config_aniversario(), ok=True, mensagem="Aniversário salvo.")


def _faz_aniversario(dia, mes, hoje):
    if not dia or not mes:
        return False
    if (dia, mes) == (hoje.day, hoje.month):
        return True
    bissexto = hoje.year % 4 == 0 and (hoje.year % 100 != 0 or hoje.year % 400 == 0)
    return (dia, mes) == (29, 2) and not bissexto and (hoje.day, hoje.month) == (28, 2)


def aniversariantes(hoje=None, cur=None):
    hoje = hoje or date.today()
    garantir_tabelas()
    proprio = cur is None
    conn = _conexao() if proprio else None
    try:
        cur = cur or conn.cursor()
        cur.execute("SELECT ClienteID, Nome, Telefone, AniversarioDia, AniversarioMes, Aceite, SaiuEm, AniversarioEnviadoAno FROM Clientes")
        return [(r[0], r[1], r[2]) for r in cur.fetchall()
                if _recebe(r[5], r[6]) and _faz_aniversario(r[3], r[4], hoje) and r[7] != hoje.year]
    finally:
        if proprio:
            conn.close()


def criar_campanha_aniversario(hoje=None):
    """10h: os aniversariantes de hoje entram numa mensagem só (fila do robô). Devolve quantos."""
    hoje = hoje or date.today()
    cfg = config_aniversario()
    if not cfg['ligado'] or not configurado():
        return 0
    foto = cfg['foto'] or None
    if foto:
        try:
            caminho_foto(foto)
        except ErroClientes:
            foto = None
    agora = datetime.now().replace(microsecond=0)
    conn = _conexao()
    try:
        cur = conn.cursor()
        lista = aniversariantes(hoje, cur)
        if not lista:
            return 0
        _criar(cur, 'aniversario', f"🎂 Aniversariantes de {hoje:%d/%m}", cfg['texto'], foto, "Aniversariantes do dia",
               lista, ST_ENVIANDO, None, 'robô', agora)
        for cid, _n, _t in lista:
            cur.execute("UPDATE Clientes SET AniversarioEnviadoAno = ? WHERE ClienteID = ?", hoje.year, cid)
        conn.commit()
    finally:
        conn.close()
    logger.info(f"[CLIENTES] Aniversário: {len(lista)} cliente(s) na fila.")
    return len(lista)


# ------------------------------------------------------------------------------
# "SAIR" respondido no WhatsApp (webhook da Z-API: "Ao receber")
# ------------------------------------------------------------------------------
def webhook_segredo():
    s = _params('webhook_segredo').get('webhook_segredo')
    if not s:
        s = secrets.token_urlsafe(18)
        _salvar_params({'webhook_segredo': s})
    return s


def _palavra(texto):
    return re.sub(r'[^a-z ]', '', _sem_acento(str(texto or '')).lower()).strip()


def processar_webhook(segredo, payload, em_segundo_plano=True):
    """A Z-API avisa cada mensagem recebida no número dos clientes. Só 'SAIR'/'VOLTAR' (a mensagem inteira) fazem algo."""
    certo = _params('webhook_segredo').get('webhook_segredo') or ''
    if not certo or not hmac.compare_digest(certo, str(segredo or '')):
        raise ErroClientes("Não encontrado.", 404)
    p = payload if isinstance(payload, dict) else {}
    if p.get('fromMe') or p.get('isGroup') or p.get('isNewsletter') or p.get('broadcast'):
        return {'ok': True, 'acao': 'ignorada'}
    inst = _numero()['zapi_instancia']
    if p.get('instanceId') and inst and str(p.get('instanceId')) != inst:
        return {'ok': True, 'acao': 'ignorada'}
    texto = (p.get('text') or {}).get('message') if isinstance(p.get('text'), dict) else None
    palavra = _palavra(texto)
    if palavra not in PALAVRAS_SAIR and palavra not in PALAVRAS_VOLTAR:
        return {'ok': True, 'acao': 'ignorada'}
    try:
        tel = normalizar_telefone(p.get('phone'))
    except ErroClientes:
        return {'ok': True, 'acao': 'ignorada'}
    agora = datetime.now().replace(microsecond=0)
    conn = _conexao()
    try:
        cur = conn.cursor()
        r = _por_telefone(cur, tel)
        if not r:
            return {'ok': True, 'acao': 'desconhecido'}
        cid, nome, aceite, saiu = r
        if palavra in PALAVRAS_SAIR:
            cur.execute("UPDATE Clientes SET SaiuEm = ?, AtualizadoEm = ? WHERE ClienteID = ?", agora, agora, cid)
            resposta, acao = TEXTO_SAIU, 'saiu'
        else:
            if not saiu and aceite:
                return {'ok': True, 'acao': 'ja_recebe'}
            cur.execute("UPDATE Clientes SET SaiuEm = NULL, Aceite = 1, AceiteEm = ?, AtualizadoEm = ? WHERE ClienteID = ?",
                        agora, agora, cid)
            resposta, acao = TEXTO_VOLTOU, 'voltou'
        conn.commit()
    finally:
        conn.close()
    logger.info(f"[CLIENTES] {nome} respondeu {palavra.upper()} no WhatsApp: {acao}.")
    (_em_segundo_plano if em_segundo_plano else (lambda f, *a: f(*a)))(enviar_para, tel, resposta)
    return {'ok': True, 'acao': acao}


# ------------------------------------------------------------------------------
# 📱 QR do balcão: o cliente se cadastra sozinho (página pública /clientes/cadastro)
# ------------------------------------------------------------------------------
def _limite_ip(ip):
    agora = time.time()
    with _trava_cadastro:
        hoje = date.today()
        if _cadastros_do_dia['dia'] != hoje:
            _cadastros_do_dia.update(dia=hoje, n=0)
        if _cadastros_do_dia['n'] >= CADASTROS_POR_DIA:
            return False
        tentativas = [t for t in _cadastros_por_ip.get(ip, []) if agora - t < CADASTROS_POR_IP[1]]
        if len(tentativas) >= CADASTROS_POR_IP[0]:
            _cadastros_por_ip[ip] = tentativas
            return False
        tentativas.append(agora)
        _cadastros_por_ip[ip] = tentativas
        _cadastros_do_dia['n'] += 1
        if len(_cadastros_por_ip) > 5000:
            _cadastros_por_ip.clear()
        return True


def cadastro_publico(dados, ip, em_segundo_plano=True):
    """O cliente se cadastra pelo QR (já autorizando). Telefone que já existe: atualiza e volta a receber."""
    dados = dados if isinstance(dados, dict) else {}
    if dados.get('site'):                     # campo escondido: só robô de spam preenche
        return {'ok': True, 'nome': ''}
    nome = _limpo(dados.get('nome'), 80)
    if len(nome) < 2:
        raise ErroClientes("Digite o seu nome.")
    tel = normalizar_telefone(dados.get('telefone'))
    dia = mes = None
    if dados.get('dia') or dados.get('mes'):
        dia, mes = _aniversario(f"{dados.get('dia')}/{dados.get('mes')}")
    if not dados.get('aceite'):
        raise ErroClientes("Marque que você aceita receber as mensagens do Gela Boca.")
    if not _limite_ip(ip or '?'):
        logger.warning(f"[CLIENTES] Cadastro pelo QR recusado (limite): IP {ip}.")
        raise ErroClientes("Muitos cadastros agora. Tente de novo mais tarde ou peça ajuda no balcão.", 429)
    garantir_tabelas()
    agora = datetime.now().replace(microsecond=0)
    conn = _conexao()
    try:
        cur = conn.cursor()
        r = _por_telefone(cur, tel)
        if r:
            cid, voltou = r[0], bool(r[3]) or not r[2]
            cur.execute("UPDATE Clientes SET Nome = ?, AniversarioDia = COALESCE(?, AniversarioDia), "
                        "AniversarioMes = COALESCE(?, AniversarioMes), Aceite = 1, AceiteEm = ?, SaiuEm = NULL, AtualizadoEm = ? "
                        "WHERE ClienteID = ?", nome, dia, mes, agora, agora, cid)
        else:
            voltou = True
            cur.execute("INSERT INTO Clientes (Nome, Telefone, AniversarioDia, AniversarioMes, Aceite, AceiteEm, Origem, CriadoEm, "
                        "AtualizadoEm) VALUES (?, ?, ?, ?, 1, ?, 'qr', ?, ?); SELECT SCOPE_IDENTITY();",
                        nome, tel, dia, mes, agora, agora, agora)
            cur.nextset()
            cid = int(cur.fetchone()[0])
        _gravar_etiquetas(cur, cid, [ETIQUETA_BALCAO], somar=True)
        conn.commit()
    finally:
        conn.close()
    logger.info(f"[CLIENTES] Cadastro pelo QR: cliente {cid}.")
    if voltou and configurado():
        (_em_segundo_plano if em_segundo_plano else (lambda f, *a: f(*a)))(
            enviar_para, tel, personalizar(TEXTO_BOAS_VINDAS, nome))
    return {'ok': True, 'nome': primeiro_nome(nome)}


def endereco_cadastro(host_url):
    """Endereço que vai no QR: o gravado em ⚙️, senão o desta página (avisa se é só da rede da loja)."""
    base = (_params('url_publica').get('url_publica') or '').strip() or str(host_url or '')
    url = base.rstrip('/') + '/clientes/cadastro'
    host = re.sub(r'^https?://', '', base).split('/')[0].split(':')[0]
    local = host in ('localhost', '127.0.0.1') or bool(re.fullmatch(r'[\d.]+', host)) or host.endswith('.local')
    return url, local


def salvar_endereco(dados, usuario):
    url = _limpo(dados.get('url'), 200).replace(' ', '')
    if url and not re.fullmatch(r'https?://[A-Za-z0-9.-]+(:\d+)?/?', url):
        raise ErroClientes("Endereço inválido. Ex.: https://compras.seusite.com.br")
    _salvar_params({'url_publica': url.rstrip('/') or None})
    logger.info(f"[CLIENTES] {usuario.get('nome')}: endereço do QR = {url or '(o da página)'}.")
    return {'ok': True, 'mensagem': "Endereço do QR salvo."}


def qr_svg(url):
    try:
        import qrcode
        import qrcode.image.svg
    except ImportError:
        return None
    img = qrcode.make(url, image_factory=qrcode.image.svg.SvgPathImage, box_size=10, border=2)
    return img.to_string(encoding='unicode')


def ajustes(host_url):
    url, local = endereco_cadastro(host_url)
    return {'numero': ajustes_numero(), 'aniversario': config_aniversario(),
            'hoje_aniversariantes': len(aniversariantes()),
            'webhook': str(host_url or '').rstrip('/') + f"/webhook/zapi-clientes/{webhook_segredo()}",
            'qr': {'url': url, 'local': local, 'svg': qr_svg(url), 'url_gravada': _params('url_publica').get('url_publica') or ''},
            'modelos': [{'nome': n, 'texto': t} for n, t in MODELOS],
            'regras': {'intervalo': f"{INTERVALO_SEG[0]} a {INTERVALO_SEG[1]} s", 'limite_dia': LIMITE_POR_DIA,
                       'horario': f"{HORARIO_ENVIO[0]}h às {HORARIO_ENVIO[1]}h"}}


# ------------------------------------------------------------------------------
# 📥 Importar: planilha (Excel/CSV) ou agenda do celular
# ------------------------------------------------------------------------------
NOMES_COLUNA = {
    'nome': ('nome', 'cliente', 'name'),
    'telefone': ('telefone', 'whatsapp', 'celular', 'fone', 'numero', 'tel', 'contato', 'phone'),
    'aniversario': ('aniversario', 'nascimento', 'data de nascimento', 'niver', 'data'),
    'etiquetas': ('etiqueta', 'etiquetas', 'grupo', 'grupos', 'tag', 'tags'),
}


def _linhas_planilha(nome_arquivo, conteudo):
    nome_arquivo = str(nome_arquivo or '').lower()
    if nome_arquivo.endswith('.xlsx'):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        except Exception:
            raise ErroClientes("Não consegui abrir a planilha. Salve como .xlsx (Excel) ou .csv e tente de novo.")
        ws = wb.worksheets[0]
        linhas = [list(r) for r, _ in zip(ws.iter_rows(values_only=True), range(MAX_LINHAS_IMPORTAR + 2))]
        wb.close()
        return linhas
    if nome_arquivo.endswith('.csv') or nome_arquivo.endswith('.txt'):
        for cod in ('utf-8-sig', 'latin-1'):
            try:
                texto = conteudo.decode(cod)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialeto = csv.Sniffer().sniff(texto[:2000], delimiters=';,\t')
        except csv.Error:
            dialeto = csv.excel
        return [r for r, _ in zip(csv.reader(io.StringIO(texto), dialeto), range(MAX_LINHAS_IMPORTAR + 2))]
    raise ErroClientes("Use uma planilha .xlsx (Excel) ou .csv.")


def _colunas(cabecalho):
    pos = {}
    for i, c in enumerate(cabecalho):
        t = _sem_acento(str(c or '')).lower().strip()
        for campo, nomes in NOMES_COLUNA.items():
            if campo not in pos and any(t == n or t.startswith(n) for n in nomes):
                pos[campo] = i
    return pos


def _ler_importacao(dados):
    """[{'nome','telefone','aniversario','etiquetas'}] da planilha (base64) ou da agenda do celular."""
    if dados.get('contatos') is not None:
        contatos = dados.get('contatos')
        if not isinstance(contatos, list) or len(contatos) > MAX_LINHAS_IMPORTAR:
            raise ErroClientes("Contatos inválidos.")
        return [{'nome': c.get('nome'), 'telefone': c.get('telefone'), 'aniversario': None, 'etiquetas': ''}
                for c in contatos if isinstance(c, dict)], 'agenda'
    b64 = str(dados.get('conteudo') or '')
    if ',' in b64[:100]:
        b64 = b64.split(',', 1)[1]
    try:
        conteudo = base64.b64decode(b64)
    except (ValueError, TypeError):
        raise ErroClientes("Arquivo inválido. Escolha a planilha de novo.")
    if not conteudo:
        raise ErroClientes("Escolha a planilha.")
    if len(conteudo) > MAX_PLANILHA_BYTES:
        raise ErroClientes("Planilha grande demais (máximo 5 MB).")
    linhas = [r for r in _linhas_planilha(dados.get('arquivo'), conteudo) if any(str(c or '').strip() for c in r)]
    if not linhas:
        raise ErroClientes("A planilha está vazia.")
    pos = _colunas(linhas[0])
    if 'telefone' in pos:
        linhas = linhas[1:]
    else:
        pos = {'nome': 0, 'telefone': 1}             # sem cabeçalho: A = nome, B = telefone
    if len(linhas) > MAX_LINHAS_IMPORTAR:
        raise ErroClientes(f"No máximo {MAX_LINHAS_IMPORTAR} clientes por vez.")
    pega = lambda r, campo: r[pos[campo]] if campo in pos and pos[campo] < len(r) else None
    return [{'nome': pega(r, 'nome'), 'telefone': pega(r, 'telefone'), 'aniversario': pega(r, 'aniversario'),
             'etiquetas': pega(r, 'etiquetas') or ''} for r in linhas], 'planilha'


def _analisar(itens, cur):
    cur.execute("SELECT Telefone, Nome FROM Clientes")
    existentes = {t: n for t, n in cur.fetchall()}
    vistos, saida = set(), []
    for it in itens:
        linha = {'nome': _limpo(it['nome'], 120), 'telefone': '', 'telefone_fmt': str(it['telefone'] or ''),
                 'aniversario': None, 'etiquetas': [], 'situacao': 'novo', 'erro': ''}
        try:
            tel_bruto = it['telefone']
            if isinstance(tel_bruto, float):
                tel_bruto = f"{tel_bruto:.0f}"         # o Excel guarda o número como 6.69999e10
            linha['telefone'] = normalizar_telefone(tel_bruto)
            linha['telefone_fmt'] = fmt_telefone(linha['telefone'])
            dia, mes = _aniversario(it['aniversario'])
            linha['aniversario'] = f"{dia:02d}/{mes:02d}" if dia else None
            linha['etiquetas'] = _etiquetas(it['etiquetas'])
            if len(linha['nome']) < 2:
                linha['nome'] = ''
            if linha['telefone'] in vistos:
                linha['situacao'], linha['erro'] = 'repetido', 'repetido na lista'
            elif linha['telefone'] in existentes:
                linha['situacao'], linha['erro'] = 'existe', f"já cadastrado ({existentes[linha['telefone']]})"
            vistos.add(linha['telefone'])
        except ErroClientes as e:
            linha['situacao'], linha['erro'] = 'invalido', str(e)
        saida.append(linha)
    return saida


def previa_importacao(dados):
    itens, origem = _ler_importacao(dados)
    garantir_tabelas()
    conn = _conexao()
    try:
        linhas = _analisar(itens, conn.cursor())
    finally:
        conn.close()
    conta = lambda s: sum(1 for l in linhas if l['situacao'] == s)
    return {'origem': origem, 'linhas': linhas[:300], 'total': len(linhas), 'novos': conta('novo'),
            'existentes': conta('existe'), 'invalidos': conta('invalido'), 'repetidos': conta('repetido')}


def importar(dados, usuario):
    """Cadastra os novos; quem já existe só ganha as etiquetas. 'autorizados' = eles aceitaram receber."""
    itens, origem = _ler_importacao(dados)
    etiqueta_extra = _etiquetas(dados.get('etiqueta') or '')
    autorizados = bool(dados.get('autorizados'))
    agora = datetime.now().replace(microsecond=0)
    garantir_tabelas()
    conn = _conexao()
    novos = atualizados = 0
    try:
        cur = conn.cursor()
        for l in _analisar(itens, cur):
            if l['situacao'] in ('invalido', 'repetido'):
                continue
            etqs = (l['etiquetas'] + etiqueta_extra)[:MAX_ETIQUETAS]
            if l['situacao'] == 'existe':
                cid = _por_telefone(cur, l['telefone'])[0]
                _gravar_etiquetas(cur, cid, etqs, somar=True)
                atualizados += 1
                continue
            dia, mes = _aniversario(l['aniversario']) if l['aniversario'] else (None, None)
            cur.execute("INSERT INTO Clientes (Nome, Telefone, AniversarioDia, AniversarioMes, Aceite, AceiteEm, Origem, CriadoEm, "
                        "AtualizadoEm) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?); SELECT SCOPE_IDENTITY();",
                        l['nome'] or fmt_telefone(l['telefone']), l['telefone'], dia, mes, 1 if autorizados else 0,
                        agora if autorizados else None, origem, agora, agora)
            cur.nextset()
            cid = int(cur.fetchone()[0])
            _gravar_etiquetas(cur, cid, etqs)
            novos += 1
        conn.commit()
    finally:
        conn.close()
    logger.info(f"[CLIENTES] {usuario.get('nome')}: importou {novos} novo(s) e atualizou {atualizados} ({origem}).")
    return {'ok': True, 'novos': novos, 'atualizados': atualizados,
            'mensagem': f"{novos} cliente(s) novo(s)" + (f", {atualizados} já existiam (ganharam as etiquetas)" if atualizados else '')
                        + ('.' if autorizados else '. Sem autorização marcada: eles ainda não recebem as mensagens.')}
